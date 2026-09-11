"""
CancunBot — Diálogo Modal de Análisis de PDFs de Pases de Caja (PAQ.pdf)
Muestra la barra de progreso, spinner animado, métricas KPI en tiempo real,
tabla interactiva de folios extraídos y opciones de exportación a Excel / creación de Lote.
"""
import logging
from typing import List, Dict, Any, Optional
from pathlib import Path

import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side

from PySide6.QtCore import Qt, QTimer, QRectF
from PySide6.QtGui import QPainter, QColor, QPen, QFont
from PySide6.QtWidgets import (
    QDialog, QVBoxLayout, QHBoxLayout, QLabel, QPushButton, QProgressBar,
    QTableWidget, QTableWidgetItem, QHeaderView, QFrame, QMessageBox, QFileDialog, QWidget, QLineEdit, QComboBox
)

from sar.src.ui.design_system.tokens.colors import Colors
from sar.src.ui.design_system.components import (
    CustomButton, CustomInput, CustomComboBox, MetricBox,
    GLMessageDialog, DialogType
)
from cancunbot.src.services.pdf_analysis_worker import PdfAnalysisWorker

logger = logging.getLogger(__name__)


class LoadingSpinner(QWidget):
    """Spinner gráfico de carga animado dinámicamente."""
    def __init__(self, color=Colors.ACCENT_BLUE, size=28, parent=None):
        super().__init__(parent)
        self.color = QColor(color)
        self.setFixedSize(size, size)
        self.angle = 0
        self.timer = QTimer(self)
        self.timer.timeout.connect(self._rotate)
        self.timer.start(25)

    def _rotate(self):
        self.angle = (self.angle + 12) % 360
        self.update()

    def paintEvent(self, event):
        painter = QPainter(self)
        painter.setRenderHint(QPainter.Antialiasing)
        pen = QPen(self.color, 3.5)
        pen.setCapStyle(Qt.RoundCap)
        painter.setPen(pen)
        rect = QRectF(3, 3, self.width() - 6, self.height() - 6)
        painter.drawArc(rect, -self.angle * 16, 270 * 16)


class KpiCardWidget(QFrame):
    """Tarjeta visual KPI de métricas en tiempo real."""
    def __init__(self, title: str, value: str = "0", icon: str = "📊", accent_color: str = Colors.ACCENT_BLUE, parent=None):
        super().__init__(parent)
        self.setStyleSheet(f"""
            QFrame {{
                background-color: {Colors.SURFACE_DARK};
                border: 1px solid {Colors.BORDER_DARK};
                border-radius: 10px;
                padding: 10px;
            }}
        """)
        layout = QHBoxLayout(self)
        layout.setContentsMargins(12, 10, 12, 10)

        # Icon Label
        lbl_icon = QLabel(icon)
        lbl_icon.setFont(QFont("Segoe UI Emoji", 20))
        layout.addWidget(lbl_icon)

        # Text Layout
        vbox = QVBoxLayout()
        vbox.setSpacing(2)
        
        lbl_title = QLabel(title.upper())
        lbl_title.setStyleSheet(f"color: {Colors.TEXT_DARK_MUTED}; font-size: 10px; font-weight: bold;")
        vbox.addWidget(lbl_title)

        self.lbl_value = QLabel(value)
        self.lbl_value.setStyleSheet(f"color: {accent_color}; font-size: 18px; font-weight: bold;")
        vbox.addWidget(self.lbl_value)

        layout.addLayout(vbox)

    def update_value(self, new_val: str):
        self.lbl_value.setText(new_val)


class PdfAnalysisDialog(QDialog):
    """
    Diálogo interactivo para previsualización, análisis asíncrono y confirmación de folios PDF.
    """
    def __init__(self, pdf_paths: List[str], api_client=None, db_connector=None, usuario_id: int = 1, parent=None):
        super().__init__(parent)
        self.pdf_paths = pdf_paths
        self.api_client = api_client
        self.db_connector = db_connector
        self.usuario_id = usuario_id
        
        self.extracted_records: List[Dict[str, Any]] = []
        self.is_finished = False
        self.created_lote_id: Optional[int] = None

        self.setWindowTitle("Análisis de Pases de Caja PDF — R2F Cancún")
        self.resize(1100, 680)
        self.setStyleSheet(f"background-color: {Colors.BG_LIGHT}; color: {Colors.TEXT_LIGHT_PRIMARY};")

        self._setup_ui()
        self._load_desarrollos_catalog()
        self._start_worker()

    def _setup_ui(self):
        main_layout = QVBoxLayout(self)
        main_layout.setSpacing(14)
        main_layout.setContentsMargins(18, 18, 18, 18)

        # 1. Cabecera con Spinner y Progreso
        header_frame = QFrame()
        header_frame.setStyleSheet(f"background-color: {Colors.SURFACE_LIGHT}; border-radius: 12px; border: 1px solid {Colors.BORDER_LIGHT};")
        header_layout = QHBoxLayout(header_frame)
        header_layout.setContentsMargins(16, 12, 16, 12)

        self.spinner = LoadingSpinner(color=Colors.ACCENT_BLUE, size=32)
        header_layout.addWidget(self.spinner)

        info_vbox = QVBoxLayout()
        info_vbox.setSpacing(4)
        
        self.lbl_header_title = QLabel("Procesando y Analizando Archivos PDF...")
        self.lbl_header_title.setStyleSheet(f"font-size: 15px; font-weight: bold; color: {Colors.TEXT_LIGHT_PRIMARY};")
        info_vbox.addWidget(self.lbl_header_title)

        self.lbl_status_subtitle = QLabel("Iniciando motor de lectura...")
        self.lbl_status_subtitle.setStyleSheet(f"color: {Colors.TEXT_LIGHT_MUTED}; font-size: 12px;")
        info_vbox.addWidget(self.lbl_status_subtitle)

        self.progress_bar = QProgressBar()
        self.progress_bar.setStyleSheet(f"""
            QProgressBar {{
                background-color: {Colors.SLATE_100};
                border: 1px solid {Colors.BORDER_LIGHT};
                border-radius: 6px;
                height: 12px;
                text-align: center;
                color: transparent;
            }}
            QProgressBar::chunk {{
                background-color: {Colors.ACCENT_BLUE};
                border-radius: 5px;
            }}
        """)
        self.progress_bar.setValue(0)
        info_vbox.addWidget(self.progress_bar)

        header_layout.addLayout(info_vbox, stretch=1)
        main_layout.addWidget(header_frame)

        # 2. Panel de Tarjetas KPI utilizando la molécula MetricBox de SAR Design System
        kpi_layout = QHBoxLayout()
        kpi_layout.setSpacing(12)

        self.card_paginas = MetricBox("Páginas Leídas", "0", Colors.ACCENT_BLUE)
        self.card_validos = MetricBox("Folios Válidos", "0", Colors.ACCENT_EMERALD)
        self.card_monto = MetricBox("Monto Total", "$ 0.00", Colors.WARNING)
        self.card_alertas = MetricBox("Advertencias", "0", Colors.ERROR)

        kpi_layout.addWidget(self.card_paginas)
        kpi_layout.addWidget(self.card_validos)
        kpi_layout.addWidget(self.card_monto)
        kpi_layout.addWidget(self.card_alertas)
        main_layout.addLayout(kpi_layout)

        # 3. Tabla Resumen en Vivo
        self.table = QTableWidget()
        self.table.setColumnCount(11)
        self.table.setHorizontalHeaderLabels([
            "#", "Archivo PDF", "Pág", "Folio Pase Caja", "RFC",
            "Nombre Contribuyente", "Emisión", "Padrón", "Clave Catastral", "Total ($)", "Estado"
        ])
        self.table.setStyleSheet(f"""
            QTableWidget {{
                background-color: {Colors.SURFACE_LIGHT};
                color: {Colors.TEXT_LIGHT_PRIMARY};
                gridline-color: {Colors.SLATE_200};
                border: 1px solid {Colors.BORDER_LIGHT};
                border-radius: 8px;
                font-size: 11px;
            }}
            QHeaderView::section {{
                background-color: {Colors.SLATE_100};
                color: {Colors.TEXT_LIGHT_PRIMARY};
                font-weight: bold;
                padding: 6px;
                border: 1px solid {Colors.SLATE_200};
            }}
        """)
        self.table.horizontalHeader().setSectionResizeMode(QHeaderView.Interactive)
        self.table.horizontalHeader().setStretchLastSection(True)
        self.table.horizontalHeader().resizeSection(0, 40)
        self.table.horizontalHeader().resizeSection(1, 140)
        self.table.horizontalHeader().resizeSection(2, 45)
        self.table.horizontalHeader().resizeSection(3, 130)
        self.table.horizontalHeader().resizeSection(4, 110)
        self.table.horizontalHeader().resizeSection(5, 180)
        self.table.horizontalHeader().resizeSection(9, 90)

        main_layout.addWidget(self.table, stretch=1)

        # 4. Campos de Entrada (Átomos CustomInput y CustomComboBox)
        metadata_row = QHBoxLayout()
        metadata_row.setSpacing(12)

        # CustomInput Descripción
        lbl_desc = QLabel("📝 Descripción:")
        lbl_desc.setStyleSheet(f"color: {Colors.TEXT_LIGHT_PRIMARY}; font-weight: bold; font-size: 12px;")
        metadata_row.addWidget(lbl_desc)

        default_desc = f"Importación PDF ({Path(self.pdf_paths[0]).name})" if self.pdf_paths else "Importación PDF Pases de Caja"
        self.txt_descripcion = CustomInput(default_desc)
        self.txt_descripcion.setPlaceholderText("Descripción personalizada del lote...")
        metadata_row.addWidget(self.txt_descripcion, stretch=3)

        # CustomComboBox Desarrollo
        lbl_desarrollo = QLabel("🏢 Desarrollo:")
        lbl_desarrollo.setStyleSheet(f"color: {Colors.TEXT_LIGHT_PRIMARY}; font-weight: bold; font-size: 12px;")
        metadata_row.addWidget(lbl_desarrollo)

        self.cb_desarrollo = CustomComboBox()
        self.cb_desarrollo.addItem("-- Sin Desarrollo Asignado --", None)
        metadata_row.addWidget(self.cb_desarrollo, stretch=2)

        main_layout.addLayout(metadata_row)

        # 5. Barra de Botones de Acción usando el átomo CustomButton de SAR
        actions_layout = QHBoxLayout()
        actions_layout.setSpacing(12)

        self.btn_export_excel = CustomButton("📥 Descargar Plantilla Excel (.xlsx)", is_secondary=True)
        self.btn_export_excel.clicked.connect(self._export_to_excel)
        actions_layout.addWidget(self.btn_export_excel)

        actions_layout.addStretch()

        self.btn_cancel = CustomButton("❌ Cancelar", is_secondary=True)
        self.btn_cancel.clicked.connect(self._on_cancel_clicked)
        actions_layout.addWidget(self.btn_cancel)

        self.btn_confirm = CustomButton("✅ Confirmar y Crear Lote", is_secondary=False)
        self.btn_confirm.setStyleSheet(f"""
            QPushButton {{
                background-color: {Colors.ACCENT_EMERALD};
                color: white;
                font-weight: bold;
                padding: 10px 20px;
                border-radius: 8px;
                border: none;
            }}
            QPushButton:hover {{ background-color: #059669; }}
            QPushButton:disabled {{ background-color: {Colors.BORDER_DARK}; color: {Colors.TEXT_DARK_MUTED}; }}
        """)
        self.btn_confirm.setEnabled(False)
        self.btn_confirm.clicked.connect(self._on_confirm_clicked)
        actions_layout.addWidget(self.btn_confirm)

        main_layout.addLayout(actions_layout)

    def _start_worker(self):
        self.worker = PdfAnalysisWorker(self.pdf_paths)
        self.worker.progress_updated.connect(self._on_progress_updated)
        self.worker.page_extracted.connect(self._on_page_extracted)
        self.worker.metric_updated.connect(self._on_metric_updated)
        self.worker.finished_analysis.connect(self._on_finished_analysis)
        self.worker.start()

    def _on_progress_updated(self, current: int, total: int, filename: str):
        self.progress_bar.setMaximum(total)
        self.progress_bar.setValue(current)
        self.lbl_status_subtitle.setText(f"Analizando {filename} — Página {current} de {total}")

    def _on_page_extracted(self, data: dict):
        self.extracted_records.append(data)
        row_idx = self.table.rowCount()
        self.table.insertRow(row_idx)

        def make_item(val: str, align_right: bool = False):
            item = QTableWidgetItem(str(val if val is not None else ""))
            if align_right:
                item.setTextAlignment(Qt.AlignRight | Qt.AlignVCenter)
            return item

        self.table.setItem(row_idx, 0, make_item(str(row_idx + 1)))
        self.table.setItem(row_idx, 1, make_item(data.get("archivo_pdf", "")))
        self.table.setItem(row_idx, 2, make_item(str(data.get("pagina", ""))))
        self.table.setItem(row_idx, 3, make_item(data.get("folio_pase_caja", "-")))
        self.table.setItem(row_idx, 4, make_item(data.get("rfc", "-")))
        self.table.setItem(row_idx, 5, make_item(data.get("nombre_contribuyente", "-")))
        self.table.setItem(row_idx, 6, make_item(data.get("fecha_emision", "-")))
        self.table.setItem(row_idx, 7, make_item(data.get("padron", "-")))
        self.table.setItem(row_idx, 8, make_item(data.get("clave_catastral", "-")))
        
        tot = data.get("total", 0.0)
        self.table.setItem(row_idx, 9, make_item(f"${tot:,.2f}", align_right=True))

        es_valido = data.get("es_valido", True)
        obs = data.get("observacion", "OK")
        item_st = make_item(obs)
        if es_valido:
            item_st.setForeground(QColor(Colors.ACCENT_EMERALD))
        else:
            item_st.setForeground(QColor(Colors.ERROR))
        self.table.setItem(row_idx, 10, item_st)

        self.table.scrollToBottom()

    def _on_metric_updated(self, name: str, val: Any):
        if name == "total_paginas":
            self.card_paginas.set_value(str(val))
        elif name == "validos":
            self.card_validos.set_value(str(val))
        elif name == "alertas":
            self.card_alertas.set_value(str(val))
        elif name == "monto_total":
            self.card_monto.set_value(f"$ {val:,.2f}")

    def _on_finished_analysis(self, success: bool, message: str, all_records: list):
        self.is_finished = True
        self.spinner.hide()
        self.progress_bar.setValue(self.progress_bar.maximum())

        if success:
            self.lbl_header_title.setText("✅ Análisis Completado Exitosamente")
            self.lbl_header_title.setStyleSheet(f"font-size: 15px; font-weight: bold; color: {Colors.ACCENT_EMERALD};")
            self.lbl_status_subtitle.setText(message)
            self.btn_confirm.setEnabled(len(self.extracted_records) > 0)
        else:
            self.lbl_header_title.setText("⚠️ Análisis Finalizado con Inconvenientes")
            self.lbl_header_title.setStyleSheet(f"font-size: 15px; font-weight: bold; color: {Colors.WARNING};")
            self.lbl_status_subtitle.setText(message)

    def _export_to_excel(self):
        if not self.extracted_records:
            QMessageBox.warning(self, "Sin datos", "No hay folios extraídos para exportar.")
            return

        file_path, _ = QFileDialog.getSaveFileName(
            self,
            "Guardar Plantilla Excel con Folios Extraídos",
            "Folios_Extraidos_Cancun.xlsx",
            "Archivos Excel (*.xlsx)"
        )
        if not file_path:
            return

        try:
            wb = openpyxl.Workbook()
            ws = wb.active
            ws.title = "Folios PAQ Cancún"

            headers = [
                "FOLIO_PASE_CAJA", "FOLIO_ELECTRONICO", "TIPO_FOLIO", "RFC",
                "NOMBRE_CONTRIBUYENTE", "FECHA_EMISION", "PADRON", "CLAVE_CATASTRAL",
                "TOTAL", "DOMICILIO", "ARCHIVO_ORIGEN", "PAGINA", "ESTADO"
            ]
            ws.append(headers)

            header_fill = PatternFill(start_color="1E293B", end_color="1E293B", fill_type="solid")
            header_font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")

            for col_num, _ in enumerate(headers, 1):
                cell = ws.cell(row=1, column=col_num)
                cell.fill = header_fill
                cell.font = header_font
                cell.alignment = Alignment(horizontal="center", vertical="center")

            for r in self.extracted_records:
                tipo = "PASE_CAJA" if r.get("folio_pase_caja") else "ELECTRONICO"
                ws.append([
                    r.get("folio_pase_caja", ""),
                    r.get("folio_electronico", ""),
                    tipo,
                    r.get("rfc", ""),
                    r.get("nombre_contribuyente", ""),
                    r.get("fecha_emision", ""),
                    r.get("padron", ""),
                    r.get("clave_catastral", ""),
                    r.get("total", 0.0),
                    r.get("domicilio", ""),
                    r.get("archivo_pdf", ""),
                    r.get("pagina", 1),
                    r.get("observacion", "OK")
                ])

            wb.save(file_path)
            QMessageBox.information(self, "Exportación Exitosa", f"Se ha guardado la plantilla Excel en:\n{file_path}")
        except Exception as e:
            logger.error(f"Error exportando a Excel: {e}")
            QMessageBox.critical(self, "Error de Exportación", f"No se pudo guardar el archivo Excel: {e}")

    def _stop_worker_safely(self):
        if hasattr(self, 'worker') and self.worker is not None:
            if self.worker.isRunning():
                self.worker.stop()
                self.worker.wait(3000)

    def closeEvent(self, event):
        self._stop_worker_safely()
        super().closeEvent(event)

    def _on_cancel_clicked(self):
        self._stop_worker_safely()
        self.reject()

    def _load_desarrollos_catalog(self):
        """Carga los desarrollos inmobiliarios activos desde la BD o API."""
        try:
            use_api = (self.api_client is not None and getattr(self.api_client, "connect_via_api", False))
            if use_api:
                rows = self.api_client.request("GET", "/api/docs/cancun/desarrollos") or []
            else:
                with self.db_connector.get_session() as session:
                    from sqlalchemy import text
                    db_rows = session.execute(
                        text("SELECT desarrollo_id, nombre FROM sar_catalogo.desarrollo WHERE activo = true ORDER BY nombre ASC")
                    ).fetchall()
                    rows = [{"desarrollo_id": r[0], "nombre": r[1]} for r in db_rows]

            if rows:
                for item in rows:
                    self.cb_desarrollo.addItem(item["nombre"], item["desarrollo_id"])
        except Exception as e:
            logger.error(f"Error cargando catálogo de desarrollos: {e}")

    def _on_confirm_clicked(self):
        valid_items = [r for r in self.extracted_records if r.get("es_valido") and (r.get("folio_pase_caja") or r.get("folio_electronico"))]
        if not valid_items:
            QMessageBox.warning(self, "Sin Folios Válidos", "No hay folios válidos seleccionados para crear el lote.")
            return

        use_api = (self.api_client is not None and getattr(self.api_client, "connect_via_api", False))
        user_desc = self.txt_descripcion.text().strip() or f"Importación PDF ({Path(self.pdf_paths[0]).name})"
        selected_desarrollo_id = self.cb_desarrollo.currentData()
        selected_desarrollo_name = self.cb_desarrollo.currentText()

        # Cuadro de confirmación antes de crear el lote
        confirm_dialog = GLMessageDialog(
            title="Confirmar Creación de Lote",
            message=f"¿Deseas confirmar la creación del lote con los siguientes datos?\n\n"
                    f"• Total Folios Válidos: {len(valid_items)}\n"
                    f"• Descripción: {user_desc}\n"
                    f"• Desarrollo: {selected_desarrollo_name}",
            dialog_type=DialogType.QUESTION,
            confirm_text="Confirmar y Crear",
            cancel_text="Cancelar",
            parent=self
        )
        if confirm_dialog.exec() != QDialog.Accepted:
            return

        try:
            if use_api:
                # Importar vía REST API
                folios_payload = []
                for item in valid_items:
                    tipo = "PASE_CAJA" if item.get("folio_pase_caja") else "ELECTRONICO"
                    f_pase = item.get("folio_pase_caja").strip() if item.get("folio_pase_caja") else None
                    f_elec = item.get("folio_electronico").strip() if item.get("folio_electronico") else None
                    r_rfc = item.get("rfc").strip().upper() if item.get("rfc") else None
                    folios_payload.append({
                        "tipo_folio": tipo,
                        "folio_pase_caja": f_pase,
                        "folio_electronico": f_elec,
                        "rfc": r_rfc
                    })
                
                resp = self.api_client.request("POST", "/api/docs/cancun/lotes/importar-pdf", json={
                    "usuario_id": self.usuario_id,
                    "origen": "MANUAL",
                    "descripcion": user_desc,
                    "desarrollo_id": selected_desarrollo_id,
                    "folios": folios_payload
                })
                if resp and resp.get("lote_id"):
                    self.created_lote_id = resp.get("lote_id")
                    omitidos = resp.get("duplicados_omitidos", 0)
                    msg_extra = f"\n(Se omitieron {omitidos} folios por estar ya registrados o duplicados)." if omitidos > 0 else ""
                    QMessageBox.information(self, "Lote Creado", f"Se ha creado exitosamente el Lote #{self.created_lote_id} con {resp.get('total_folios')} folios.{msg_extra}")
                    self.accept()
                else:
                    QMessageBox.critical(self, "Error API", "No se pudo crear el lote a través del servidor API.")
            else:
                # Importar de forma directa en BD (modo LAN)
                from cancunbot.src.storage.cancunbot_repos import LoteFolioRepository, FolioCancunRepository
                from cancunbot.src.storage.cancunbot_models import FolioCancun
                from sqlalchemy import select

                with self.db_connector.get_session() as session:
                    from sqlalchemy import text
                    db_pases = set(session.scalars(select(FolioCancun.folio_pase_caja).where(FolioCancun.folio_pase_caja.isnot(None))).all())
                    db_elecs = set(session.scalars(select(FolioCancun.folio_electronico).where(FolioCancun.folio_electronico.isnot(None))).all())

                    # Cargar mapa de RFCs activos para resolución directa de rfc_id
                    rfc_rows = session.execute(text("SELECT UPPER(rfc), rfc_id FROM sar_catalogo.rfc WHERE activo = true")).fetchall()
                    rfc_map = {r[0]: r[1] for r in rfc_rows}

                    folios_list = []
                    duplicados_omitidos = 0
                    folios_vistos = set()

                    for item in valid_items:
                        tipo = "PASE_CAJA" if item.get("folio_pase_caja") else "ELECTRONICO"
                        raw_val = item.get("folio_pase_caja") if tipo == "PASE_CAJA" else item.get("folio_electronico")
                        val = raw_val.strip() if raw_val else None
                        
                        if not val or val in folios_vistos:
                            duplicados_omitidos += 1
                            continue

                        if tipo == "PASE_CAJA" and val in db_pases:
                            duplicados_omitidos += 1
                            continue
                        if tipo == "ELECTRONICO" and val in db_elecs:
                            duplicados_omitidos += 1
                            continue

                        folios_vistos.add(val)
                        f_pase = val if tipo == "PASE_CAJA" else None
                        f_elec = val if tipo == "ELECTRONICO" else None

                        # Validar RFC contra la base de datos (asignar rfc_id si existe, o None/null si está vacío/no catalogado)
                        raw_rfc = (item.get("rfc") or "").strip().upper()
                        rfc_id_val = rfc_map.get(raw_rfc) if raw_rfc else None

                        folios_list.append({
                            "tipo_folio": tipo,
                            "folio_pase_caja": f_pase,
                            "folio_electronico": f_elec,
                            "rfc_id": rfc_id_val,
                            "desarrollo_id": selected_desarrollo_id
                        })

                    if not folios_list:
                        QMessageBox.warning(self, "Sin Folios Nuevos", "Todos los folios del archivo PDF ya existen en la base de datos.")
                        return

                    lote_repo = LoteFolioRepository(session)
                    folio_repo = FolioCancunRepository(session)

                    lote = lote_repo.create(
                        usuario_id=self.usuario_id,
                        origen="MANUAL",
                        descripcion=user_desc
                    )
                    
                    inserted = folio_repo.create_bulk(lote.lote_id, folios_list)
                    lote_repo.update_metrics_and_status(lote.lote_id)
                    session.commit()

                    self.created_lote_id = lote.lote_id
                    msg_extra = f"\n(Se omitieron {duplicados_omitidos} folios duplicados/existentes)." if duplicados_omitidos > 0 else ""
                    QMessageBox.information(self, "Lote Creado", f"Se ha creado el Lote #{lote.lote_id} ({lote.folio_lote}) con {inserted} folios nuevos.{msg_extra}")
                    self.accept()
        except Exception as e:
            logger.error(f"Error creando lote desde PDF: {e}")
            QMessageBox.critical(self, "Error al Crear Lote", f"No se pudo crear el lote en el sistema: {e}")
