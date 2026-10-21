"""Ubicaciones Bulk Load and Management View for Special Processes."""

import os
from typing import List, Dict, Any, Optional
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side

from PySide6.QtWidgets import (
    QWidget, QVBoxLayout, QHBoxLayout, QFileDialog,
    QTableWidget, QTableWidgetItem, QHeaderView, QAbstractItemView, QLabel, QFrame
)
from PySide6.QtCore import Qt, QThread, Signal
from PySide6.QtGui import QColor
from sar.src.ui.design_system.components.molecules.gl_combo_box import CustomComboBox
from sar.src.ui.design_system.components.molecules.gl_card import CustomCard
from sar.src.ui.design_system.components.atoms.gl_button import CustomButton
from sar.src.ui.design_system.components.atoms.gl_label import CustomLabel
from sar.src.ui.design_system.components.organisms.gl_message_dialog import GLMessageBox as QMessageBox
from sar.src.ui.design_system.tokens.colors import Colors
from sar.src.ui.design_system.utils.icons import Icons


class UbicacionesImportWorker(QThread):
    """Worker thread to execute bulk locations import without blocking the UI."""
    finished = Signal(bool, dict, str)  # success, result_dict, error_msg

    def __init__(self, api_client, db_connector, desarrollo_id: int, records: List[Dict[str, Any]], usuario_id: int, sesion_id: Optional[int] = None):
        super().__init__()
        self.api_client = api_client
        self.db_connector = db_connector
        self.desarrollo_id = desarrollo_id
        self.records = records
        self.usuario_id = usuario_id
        self.sesion_id = sesion_id

    def run(self):
        try:
            if self.api_client.connect_via_api:
                payload = {
                    "desarrollo_id": self.desarrollo_id,
                    "usuario_id": self.usuario_id,
                    "sesion_id": self.sesion_id,
                    "records": self.records
                }
                res = self.api_client.request("POST", "/api/admin/ubicaciones/bulk-import", json=payload)
            else:
                from sar.src.storage.repositories import CatalogoRepository
                with self.db_connector.get_session() as session:
                    repo = CatalogoRepository(session)
                    res = repo.importar_catalogo_ubicaciones(
                        desarrollo_id=self.desarrollo_id,
                        records=self.records,
                        usuario_id=self.usuario_id,
                        sesion_id=self.sesion_id
                    )
                    session.commit()
            self.finished.emit(True, res, "")
        except Exception as e:
            self.finished.emit(False, {}, str(e))


class UbicacionesView(QWidget):
    """View to download the Excel template and import locations catalog per development."""

    def __init__(self, db_connector, current_user_id, current_sesion_id, can_edit, parent=None):
        super().__init__(parent)
        self.db_connector = db_connector
        self.current_user_id = current_user_id
        self.current_sesion_id = current_sesion_id
        self.can_edit = can_edit

        from sar.src.storage.api_client import APIClient
        self.api_client = APIClient()

        self._desarrollos_map = {}
        self._parsed_records = []
        self._worker = None

        self.layout = QVBoxLayout(self)
        self.layout.setContentsMargins(0, 0, 0, 0)
        self.layout.setSpacing(16)

        self._build_ui()

    def _build_ui(self):
        # 1. Informative description banner
        self.banner_card = CustomCard()
        self.lbl_desc = CustomLabel(
            "<b>Carga de Catálogo de Ubicaciones por Desarrollo:</b><br/>"
            "Esta herramienta administrativa permite precargar masivamente el padrón físico de lotes, manzanas y viviendas de un Desarrollo.<br/>"
            "1. <b>Selecciona</b> el Desarrollo destino sobre el cual deseas importar el catálogo.<br/>"
            "2. Presiona <b>Descargar Plantilla Excel</b> para obtener el formato oficial estandarizado (.xlsx).<br/>"
            "3. Llena la plantilla con las coordenadas catastrales o de obra (SM, Manzana, Lote, Edificio, Vivienda, No. Oficial).<br/>"
            "4. Presiona <b>Seleccionar Archivo</b> para previsualizar los registros analizados.<br/>"
            "5. Presiona <b>Ejecutar Importación</b> para insertar transaccionalmente las ubicaciones con detección automática de duplicados.",
            variant="body"
        )
        self.lbl_desc.setWordWrap(True)
        self.banner_card.layout.addWidget(self.lbl_desc)
        self.layout.addWidget(self.banner_card)

        # 2. Controls Toolbar
        controls_layout = QHBoxLayout()
        controls_layout.setSpacing(10)

        controls_layout.addWidget(CustomLabel("Desarrollo Destino:", variant="body"))
        self.combo_desarrollo = CustomComboBox()
        self.combo_desarrollo.setMinimumWidth(260)
        self.combo_desarrollo.currentIndexChanged.connect(self._on_desarrollo_changed)
        controls_layout.addWidget(self.combo_desarrollo)

        # Download Template button
        self.btn_descargar_plantilla = CustomButton("Descargar Plantilla Excel", is_secondary=True, icon_name="excel")
        self.btn_descargar_plantilla.clicked.connect(self._on_descargar_plantilla)
        controls_layout.addWidget(self.btn_descargar_plantilla)

        # Select Excel file button
        self.btn_seleccionar_archivo = CustomButton("Seleccionar Archivo", is_secondary=True, icon_name="documento_cargar")
        self.btn_seleccionar_archivo.clicked.connect(self._on_seleccionar_archivo)
        controls_layout.addWidget(self.btn_seleccionar_archivo)

        # Clear / Reset button
        self.btn_limpiar = CustomButton("Limpiar", is_secondary=True, icon_name="cancelar")
        self.btn_limpiar.clicked.connect(self._on_limpiar)
        controls_layout.addWidget(self.btn_limpiar)

        controls_layout.addStretch()

        # Execute Import button
        self.btn_ejecutar = CustomButton("Ejecutar Importación", icon_name="guardar")
        self.btn_ejecutar.clicked.connect(self._on_ejecutar_importacion)
        self.btn_ejecutar.setEnabled(False)
        controls_layout.addWidget(self.btn_ejecutar)

        self.layout.addLayout(controls_layout)

        # 3. KPI / Status summary bar
        self.kpi_card = CustomCard()
        kpi_layout = QHBoxLayout()
        kpi_layout.setContentsMargins(0, 0, 0, 0)
        kpi_layout.setSpacing(16)

        self.lbl_kpi_archivo = CustomLabel("Archivo: Ninguno seleccionado", variant="muted")
        self.lbl_kpi_total = CustomLabel("Total en Archivo: 0", variant="subheader")
        self.lbl_kpi_nuevos = CustomLabel("Nuevos: 0", variant="subheader")
        self.lbl_kpi_duplicados = CustomLabel("Duplicados: 0", variant="subheader")
        self.lbl_kpi_existentes = CustomLabel("Existentes en BD: 0", variant="subheader")

        kpi_layout.addWidget(self.lbl_kpi_archivo)
        kpi_layout.addStretch()
        kpi_layout.addWidget(self.lbl_kpi_total)
        kpi_layout.addWidget(self.lbl_kpi_nuevos)
        kpi_layout.addWidget(self.lbl_kpi_duplicados)
        kpi_layout.addWidget(self.lbl_kpi_existentes)
        self.kpi_card.layout.addLayout(kpi_layout)
        self.layout.addWidget(self.kpi_card)

        # 4. Preview / Catalog Table
        self.lbl_table_title = CustomLabel("Previsualización de Ubicaciones / Inventario del Desarrollo", variant="subheader")
        self.layout.addWidget(self.lbl_table_title)

        self.table = QTableWidget()
        self.table.setEditTriggers(QAbstractItemView.NoEditTriggers)
        self.table.setSelectionBehavior(QAbstractItemView.SelectRows)
        self.table.setSelectionMode(QAbstractItemView.SingleSelection)
        self.table.setColumnCount(8)
        self.table.setHorizontalHeaderLabels([
            "Estatus", "SM", "Manzana", "Lote", "Edificio", "Vivienda", "No. Oficial", "Descripción / Notas"
        ])
        self.table.horizontalHeader().setStretchLastSection(True)
        self.table.horizontalHeader().setSectionResizeMode(0, QHeaderView.ResizeToContents)
        self.table.verticalHeader().setVisible(False)
        self.table.setAlternatingRowColors(True)
        self.layout.addWidget(self.table)

    def refresh_data(self):
        """Loads available developments into the combo box."""
        self.combo_desarrollo.blockSignals(True)
        self.combo_desarrollo.clear()
        self.combo_desarrollo.addItem("-- Seleccione un Desarrollo --", None)
        self._desarrollos_map.clear()

        try:
            if self.api_client.connect_via_api:
                des_list = self.api_client.request("GET", "/api/admin/data/desarrollos")
            else:
                from sar.src.storage.repositories import CatalogoRepository
                with self.db_connector.get_session() as session:
                    repo = CatalogoRepository(session)
                    items = repo.get_all_desarrollos()
                    des_list = [
                        {"desarrollo_id": d.desarrollo_id, "nombre": d.nombre, "alias": d.alias, "activo": d.activo}
                        for d in items
                    ]

            for d in des_list:
                if d.get("activo", True):
                    self.combo_desarrollo.addItem(d["nombre"], d["desarrollo_id"])
                    self._desarrollos_map[d["desarrollo_id"]] = d["nombre"]
        except Exception as e:
            QMessageBox.critical(self, "Error", f"No se pudieron cargar los desarrollos: {e}")

        self.combo_desarrollo.blockSignals(False)
        self._on_desarrollo_changed()

    def _on_desarrollo_changed(self):
        """Loads existing locations from DB for the selected development."""
        desarrollo_id = self.combo_desarrollo.currentData()
        self._parsed_records = []
        self._existentes_keys = set()
        self.btn_ejecutar.setEnabled(False)
        self.lbl_kpi_archivo.setText("Archivo: Ninguno seleccionado")
        self.lbl_kpi_total.setText("Total en Archivo: 0")
        self.lbl_kpi_nuevos.setText("Nuevos: 0")
        self.lbl_kpi_duplicados.setText("Duplicados: 0")

        if not desarrollo_id:
            self.table.setRowCount(0)
            self.lbl_kpi_existentes.setText("Existentes en BD: 0")
            return

        try:
            if self.api_client.connect_via_api:
                ubis = self.api_client.request("GET", f"/api/admin/ubicaciones/{desarrollo_id}")
            else:
                from sar.src.storage.repositories import CatalogoRepository
                with self.db_connector.get_session() as session:
                    repo = CatalogoRepository(session)
                    ubis = repo.get_ubicaciones_by_desarrollo(desarrollo_id)

            # Construir set de claves existentes normalizadas (sm, mz, lt, ed, vi, no) sin espacios
            for u in ubis:
                key = (
                    "".join(str(u.get("sm") or "").split()).upper(),
                    "".join(str(u.get("mz") or "").split()).upper(),
                    "".join(str(u.get("lote") or "").split()).upper(),
                    "".join(str(u.get("edif") or "").split()).upper(),
                    "".join(str(u.get("viv") or "").split()).upper(),
                    "".join(str(u.get("no_oficial") or "").split()).upper()
                )
                self._existentes_keys.add(key)

            self.lbl_kpi_existentes.setText(f"Existentes en BD: {len(ubis)}")
            self.lbl_table_title.setText(f"Ubicaciones registradas en base de datos para: {self.combo_desarrollo.currentText()} ({len(ubis)})")

            self.table.setRowCount(len(ubis))
            for idx, u in enumerate(ubis):
                item_st = QTableWidgetItem("REGISTRADO")
                item_st.setForeground(Qt.darkGreen)
                self.table.setItem(idx, 0, item_st)
                self.table.setItem(idx, 1, QTableWidgetItem(str(u.get("sm") or "")))
                self.table.setItem(idx, 2, QTableWidgetItem(str(u.get("mz") or "")))
                self.table.setItem(idx, 3, QTableWidgetItem(str(u.get("lote") or "")))
                self.table.setItem(idx, 4, QTableWidgetItem(str(u.get("edif") or "")))
                self.table.setItem(idx, 5, QTableWidgetItem(str(u.get("viv") or "")))
                self.table.setItem(idx, 6, QTableWidgetItem(str(u.get("no_oficial") or "")))
                self.table.setItem(idx, 7, QTableWidgetItem(str(u.get("descripcion") or "")))

        except Exception as e:
            QMessageBox.critical(self, "Error", f"Error al consultar ubicaciones existentes: {e}")

    def _on_descargar_plantilla(self):
        """Generates a professional Excel template (.xlsx) with formatting and examples."""
        des_id = self.combo_desarrollo.currentData()
        des_nombre = self.combo_desarrollo.currentText() if des_id else "GENERAL"
        clean_des = "".join(c for c in des_nombre if c.isalnum() or c in (" ", "_", "-")).strip()

        default_filename = f"Plantilla_Ubicaciones_{clean_des}.xlsx"
        file_path, _ = QFileDialog.getSaveFileName(
            self, "Guardar Plantilla de Ubicaciones", default_filename, "Excel Files (*.xlsx)"
        )
        if not file_path:
            return

        try:
            wb = openpyxl.Workbook()
            ws = wb.active
            ws.title = "Ubicaciones"

            # Headers
            headers = ["SM", "MANZANA", "LOTE", "EDIFICIO", "VIVIENDA", "NO_OFICIAL", "DESCRIPCION"]
            ws.append(headers)

            # Styling
            header_fill = PatternFill(start_color="1E293B", end_color="1E293B", fill_type="solid")
            header_font = Font(name="Segoe UI", size=11, bold=True, color="FFFFFF")
            border_thin = Border(
                left=Side(style='thin', color='CBD5E1'),
                right=Side(style='thin', color='CBD5E1'),
                top=Side(style='thin', color='CBD5E1'),
                bottom=Side(style='thin', color='CBD5E1')
            )

            for col_num in range(1, len(headers) + 1):
                cell = ws.cell(row=1, column=col_num)
                cell.fill = header_fill
                cell.font = header_font
                cell.alignment = Alignment(horizontal="center", vertical="center")
                cell.border = border_thin

            # Add sample rows
            sample_rows = [
                ["15", "01", "10", "", "", "101", "Lote unifamiliar tipo"],
                ["15", "01", "11", "", "", "102", "Lote unifamiliar tipo"],
                ["15", "02", "01", "TORRE A", "101", "A-101", "Departamento primer nivel"],
                ["15", "02", "01", "TORRE A", "102", "A-102", "Departamento primer nivel"]
            ]
            for row_data in sample_rows:
                ws.append(row_data)

            # Column dimensions
            col_widths = {"A": 10, "B": 14, "C": 12, "D": 16, "E": 14, "F": 16, "G": 32}
            for col_letter, width in col_widths.items():
                ws.column_dimensions[col_letter].width = width

            wb.save(file_path)
            QMessageBox.information(
                self, "Plantilla Generada",
                f"Plantilla Excel generada con éxito en:\n{file_path}\n\n"
                "Puede llenarla en Excel con los predios oficiales y cargarla nuevamente."
            )
        except Exception as e:
            QMessageBox.critical(self, "Error", f"No se pudo generar la plantilla Excel: {e}")

    def _on_seleccionar_archivo(self):
        """Parses an Excel file and prepares records for preview and validation."""
        des_id = self.combo_desarrollo.currentData()
        if not des_id:
            QMessageBox.warning(self, "Desarrollo Requerido", "Por favor seleccione un Desarrollo destino antes de cargar el archivo.")
            return

        file_path, _ = QFileDialog.getOpenFileName(
            self, "Seleccionar Plantilla de Ubicaciones", "", "Excel Files (*.xlsx *.xls)"
        )
        if not file_path:
            return

        try:
            wb = openpyxl.load_workbook(file_path, data_only=True)
            ws = wb.active
            rows = list(ws.iter_rows(values_only=True))
            if not rows or len(rows) < 2:
                QMessageBox.warning(self, "Archivo Vacío", "El archivo seleccionado no contiene filas de datos.")
                return

            # Header mapping
            raw_headers = [str(h or "").strip().upper() for h in rows[0]]
            def get_col_idx(*names):
                for name in names:
                    n_up = name.upper()
                    if n_up in raw_headers:
                        return raw_headers.index(n_up)
                return -1

            idx_sm = get_col_idx("SM", "SUPERMANZANA", "REGION")
            idx_mz = get_col_idx("MANZANA", "MZA", "MZ")
            idx_lt = get_col_idx("LOTE", "LT")
            idx_ed = get_col_idx("EDIFICIO", "EDIF", "EXT")
            idx_vi = get_col_idx("VIVIENDA", "VIV", "INT")
            idx_no = get_col_idx("NO_OFICIAL", "NO. OFICIAL", "NO.OFICIAL", "OFICIAL", "FOLIO")
            idx_desc = get_col_idx("DESCRIPCION", "DESCRIPCIÓN", "NOTAS", "OBSERVACIONES")

            if idx_mz == -1 and idx_lt == -1 and idx_no == -1:
                QMessageBox.critical(
                    self, "Encabezados Faltantes",
                    "El archivo no contiene las columnas mínimas requeridas ('MANZANA', 'LOTE' o 'NO_OFICIAL')."
                )
                return

            parsed = []
            archivo_seen = set()
            nuevos_count = 0
            duplicados_count = 0

            for r in rows[1:]:
                if not any(r):
                    continue
                sm = "".join(str(r[idx_sm] if idx_sm != -1 and r[idx_sm] is not None else "").split()).upper()
                mz = "".join(str(r[idx_mz] if idx_mz != -1 and r[idx_mz] is not None else "").split()).upper()
                lt = "".join(str(r[idx_lt] if idx_lt != -1 and r[idx_lt] is not None else "").split()).upper()
                ed = "".join(str(r[idx_ed] if idx_ed != -1 and r[idx_ed] is not None else "").split()).upper()
                vi = "".join(str(r[idx_vi] if idx_vi != -1 and r[idx_vi] is not None else "").split()).upper()
                no = "".join(str(r[idx_no] if idx_no != -1 and r[idx_no] is not None else "").split()).upper()
                desc = str(r[idx_desc] if idx_desc != -1 and r[idx_desc] is not None else "").strip()

                if not any([mz, lt, no]):
                    continue

                # Clave de unicidad de 6 campos: (sm, mz, lt, ed, vi, no_oficial)
                coord_key = (
                    sm,
                    mz,
                    lt,
                    ed,
                    vi,
                    no
                )

                es_duplicado_bd = coord_key in self._existentes_keys
                es_duplicado_archivo = coord_key in archivo_seen

                if es_duplicado_bd:
                    estatus = "DUPLICADO EN BD"
                    duplicados_count += 1
                elif es_duplicado_archivo:
                    estatus = "DUPLICADO EN ARCHIVO"
                    duplicados_count += 1
                else:
                    estatus = "NUEVO"
                    nuevos_count += 1
                    archivo_seen.add(coord_key)

                parsed.append({
                    "sm": sm,
                    "mz": mz,
                    "lote": lt,
                    "edif": ed,
                    "viv": vi,
                    "no_oficial": no,
                    "descripcion": desc,
                    "estatus": estatus,
                    "es_duplicado": (es_duplicado_bd or es_duplicado_archivo)
                })

            if not parsed:
                QMessageBox.warning(self, "Sin Registros Válidos", "No se encontraron filas con coordenadas válidas para importar.")
                return

            self._parsed_records = parsed
            self.lbl_kpi_archivo.setText(f"Archivo: {os.path.basename(file_path)}")
            self.lbl_kpi_total.setText(f"Total en Archivo: {len(parsed)}")
            self.lbl_kpi_nuevos.setText(f"Nuevos a Insertar: {nuevos_count}")
            self.lbl_kpi_duplicados.setText(f"Duplicados Omitidos: {duplicados_count}")
            self.lbl_table_title.setText(f"Previsualización del archivo: {nuevos_count} nuevos, {duplicados_count} duplicados ({len(parsed)} total)")

            # Populate table preview
            self.table.setRowCount(len(parsed))
            for idx, p in enumerate(parsed):
                item_st = QTableWidgetItem(p["estatus"])
                if p["es_duplicado"]:
                    item_st.setForeground(QColor("#E11D48"))  # Rojo advertencia
                else:
                    item_st.setForeground(QColor("#0284C7"))  # Azul nuevo por importar
                self.table.setItem(idx, 0, item_st)
                self.table.setItem(idx, 1, QTableWidgetItem(p["sm"]))
                self.table.setItem(idx, 2, QTableWidgetItem(p["mz"]))
                self.table.setItem(idx, 3, QTableWidgetItem(p["lote"]))
                self.table.setItem(idx, 4, QTableWidgetItem(p["edif"]))
                self.table.setItem(idx, 5, QTableWidgetItem(p["viv"]))
                self.table.setItem(idx, 6, QTableWidgetItem(p["no_oficial"]))
                self.table.setItem(idx, 7, QTableWidgetItem(p["descripcion"]))

            # Habilitar ejecución solo si hay registros nuevos por importar
            self.btn_ejecutar.setEnabled(self.can_edit and nuevos_count > 0)

            if duplicados_count > 0:
                QMessageBox.information(
                    self, "Validación de Duplicados",
                    f"Se analizaron {len(parsed)} registros:\n\n"
                    f"• {nuevos_count} registros nuevos listos para importar.\n"
                    f"• {duplicados_count} registros duplicados detectados (identificados con estatus DUPLICADO en color rojo y que se omitirán)."
                )

        except Exception as e:
            QMessageBox.critical(self, "Error al Leer Excel", f"Ocurrió un error al procesar el archivo:\n{e}")

    def _on_limpiar(self):
        """Limpia el archivo seleccionado, resetea el desarrollo y vacía la tabla y KPIs."""
        self._parsed_records = []
        self._existentes_keys = set()
        self.btn_ejecutar.setEnabled(False)
        self.lbl_kpi_archivo.setText("Archivo: Ninguno seleccionado")
        self.lbl_kpi_total.setText("Total en Archivo: 0")
        self.lbl_kpi_nuevos.setText("Nuevos: 0")
        self.lbl_kpi_duplicados.setText("Duplicados: 0")
        self.lbl_kpi_existentes.setText("Existentes en BD: 0")
        self.lbl_table_title.setText("Previsualización de Ubicaciones / Inventario del Desarrollo")
        self.table.setRowCount(0)
        # Resetear selector de desarrollo al elemento inicial
        self.combo_desarrollo.setCurrentIndex(0)

    def _on_ejecutar_importacion(self):
        """Launches background worker to commit locations into the database."""
        des_id = self.combo_desarrollo.currentData()
        if not des_id or not self._parsed_records:
            return

        registros_a_importar = [r for r in self._parsed_records if not r.get("es_duplicado")]
        if not registros_a_importar:
            QMessageBox.warning(
                self, "Sin Registros Nuevos",
                "Todos los registros en el archivo ya existen en la base de datos o están duplicados. No hay nada nuevo que importar."
            )
            return

        reply = QMessageBox.question(
            self, "Confirmar Carga Masiva",
            f"¿Desea importar {len(registros_a_importar)} registros de ubicación nuevos para el desarrollo '{self.combo_desarrollo.currentText()}'?\n\n"
            f"(Se omitirán {len(self._parsed_records) - len(registros_a_importar)} registros detectados como duplicados).",
            QMessageBox.Yes | QMessageBox.No
        )
        if reply != QMessageBox.Yes:
            return

        self.btn_ejecutar.setEnabled(False)
        self.btn_seleccionar_archivo.setEnabled(False)
        self.btn_descargar_plantilla.setEnabled(False)
        self.btn_limpiar.setEnabled(False)

        self._worker = UbicacionesImportWorker(
            api_client=self.api_client,
            db_connector=self.db_connector,
            desarrollo_id=des_id,
            records=registros_a_importar,
            usuario_id=self.current_user_id or 1,
            sesion_id=self.current_sesion_id
        )
        self._worker.finished.connect(self._on_import_finished)
        self._worker.start()

    def _on_import_finished(self, success: bool, res: dict, error_msg: str):
        """Callback executed when the import worker finishes."""
        self.btn_seleccionar_archivo.setEnabled(True)
        self.btn_descargar_plantilla.setEnabled(True)
        self.btn_limpiar.setEnabled(True)

        if success:
            ins = res.get("insertadas", 0)
            act = res.get("actualizadas", 0)
            omi = res.get("omitidas", 0)
            tot = res.get("total_vigentes", 0)
            QMessageBox.information(
                self, "Importación Exitosa",
                f"✅ Catálogo de ubicaciones procesado correctamente:\n\n"
                f"• Nuevas ubicaciones insertadas: {ins}\n"
                f"• Ubicaciones existentes actualizadas: {act}\n"
                f"• Registros omitidos (sin cambios): {omi}\n"
                f"• Total de ubicaciones vigentes en el desarrollo: {tot}"
            )
            # Reload from DB
            self._on_desarrollo_changed()
        else:
            QMessageBox.critical(self, "Error en la Importación", f"No se pudo completar la carga masiva:\n{error_msg}")
            self.btn_ejecutar.setEnabled(True)
