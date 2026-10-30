import os
import re
from typing import Optional, List, Dict, Any
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

from PySide6.QtWidgets import (
    QWidget, QHBoxLayout, QVBoxLayout, QFormLayout, QFrame, QLabel, QTabWidget,
    QTableWidget, QTableWidgetItem, QHeaderView, QAbstractItemView, QFileDialog,
    QSizePolicy, QLineEdit
)
from PySide6.QtCore import Qt, QThread, Signal
from PySide6.QtGui import QColor
from sar.src.ui.design_system.tokens.colors import Colors
from sar.src.ui.design_system.utils.icons import Icons
from sar.src.ui.design_system.components.atoms.gl_label import CustomLabel
from sar.src.ui.design_system.components.atoms.gl_input import CustomInput
from sar.src.ui.design_system.components.atoms.gl_button import CustomButton
from sar.src.ui.design_system.components.atoms.gl_checkbox import CustomCheckBox
from sar.src.ui.design_system.components.molecules.gl_combo_box import CustomComboBox
from sar.src.ui.design_system.components.molecules.gl_card import CustomCard
from sar.src.ui.design_system.components.molecules.gl_loading_dialog import GLLoadingDialog
from sar.src.ui.design_system.components.organisms.gl_message_dialog import GLMessageBox as QMessageBox
from sar.src.ui.design_system.components.organisms.gl_crud_table import CrudTablePanel
from sar.src.ui.design_system.components.organisms.gl_dialog import CustomDialog
from sar.src.storage.repositories import CatalogoRepository
from satc.src.services.housing_life_cycle_service import HousingLifeCycleService


class ViviendasFileAnalyzeWorker(QThread):
    """Worker thread para leer, validar y analizar archivos Excel o Constancias RPP PDF en segundo plano."""
    finished = Signal(bool, dict, str)
    progress_updated = Signal(int, int, str)

    def __init__(self, api_client, db_connector, file_path: str, desarrollo_id: int, regimen_id: Optional[int], modo: str):
        super().__init__()
        self.api_client = api_client
        self.db_connector = db_connector
        self.file_path = file_path
        self.desarrollo_id = desarrollo_id
        self.regimen_id = regimen_id
        self.modo = modo

    def run(self):
        import time
        start_time = time.time()
        try:
            records = []
            is_pdf = self.file_path.lower().endswith(".pdf")

            if is_pdf:
                if self.modo == "CATASTRO":
                    from satc.src.services.catastro_oficio_extractor import CatastroOficioExtractor

                    def _progress_cb(curr, tot):
                        self.progress_updated.emit(curr, tot, f"Procesando página {curr} de {tot}...")

                    raw_extracted = CatastroOficioExtractor.extract_from_pdf(self.file_path, progress_callback=_progress_cb)
                    if not raw_extracted:
                        self.finished.emit(
                            False, {},
                            "No se pudieron extraer registros del Oficio de Catastro PDF.\n"
                            "Verifique que el documento contenga tablas escaneadas con Claves Catastrales legibles."
                        )
                        return
                    for r in raw_extracted:
                        records.append({
                            "interior": r.get("interior", ""),
                            "edificio": r.get("edificio", ""),
                            "nivel": "",
                            "numero_oficial": r.get("numero_oficial", ""),
                            "calle": "",
                            "destino": "HABITACIONAL",
                            "espacio": "VIVIENDA",
                            "folio_electronico": r.get("folio_electronico", ""),
                            "clave_catastral": r.get("clave_catastral", ""),
                            "fecha_alta_cc": r.get("fecha_alta_cc", ""),
                            "padron": "",
                            "descripcion": r.get("descripcion", "Extraído de Oficio de Catastro")
                        })
                elif self.modo == "RPP":
                    from satc.src.services.rpp_resolution_extractor import RppResolutionExtractor

                    def _rpp_progress_cb(curr, tot):
                        self.progress_updated.emit(curr, tot, f"Procesando página {curr} de {tot}...")

                    raw_extracted = RppResolutionExtractor.extract_from_pdf(self.file_path, progress_callback=_rpp_progress_cb)

                    # Auditoría de Consecutividad: Calcular la pendiente base (Folio - Número Oficial)
                    deltas = []
                    for r in raw_extracted:
                        try:
                            nof_val = int(r.get("numero_oficial") or 0)
                            fol_val = int(r.get("folio_electronico") or 0)
                            if nof_val > 0 and fol_val > 0:
                                deltas.append(fol_val - nof_val)
                        except Exception:
                            pass

                    base_delta = None
                    if deltas:
                        import collections
                        counter = collections.Counter(deltas)
                        base_delta, _ = counter.most_common(1)[0]

                    for r in raw_extracted:
                        pag_num = r.get("pagina", 1)
                        txt_comp = r.get("texto_completo") or f"Frac. #{r['numero_fraccion']}"
                        desc_rpp = f"Constancia RPP Pág {pag_num} | {txt_comp}"
                        # Alerta si el folio o número oficial se desvía del consecutivo matemáticamente esperado
                        if base_delta is not None:
                            try:
                                nof_val = int(r.get("numero_oficial") or 0)
                                fol_val = int(r.get("folio_electronico") or 0)
                                if nof_val > 0 and fol_val > 0:
                                    curr_delta = fol_val - nof_val
                                    if curr_delta != base_delta:
                                        folio_esperado = str(base_delta + nof_val)
                                        desc_rpp = f"⚠️ FOLIO FUERA DE CONSECUTIVO (Esperado: {folio_esperado}) | {desc_rpp}"
                            except Exception:
                                pass

                        records.append({
                            "interior": r["interior"],
                            "edificio": r["edificio"],
                            "nivel": r.get("nivel", ""),
                            "numero_oficial": r.get("numero_oficial", ""),
                            "calle": "",
                            "destino": "HABITACIONAL",
                            "espacio": "VIVIENDA",
                            "folio_electronico": r["folio_electronico"],
                            "clave_catastral": "",
                            "fecha_alta_cc": "",
                            "padron": "",
                            "descripcion": desc_rpp
                        })
                else:
                    self.finished.emit(
                        False, {},
                        "El archivo seleccionado es un PDF.\n"
                        "Por favor seleccione '2. Vincular Folios RPP' o '3. Vincular CC / Catastro'."
                    )
                    return
            else:
                self.progress_updated.emit(5, 100, "Cargando libro de Excel...")
                wb = openpyxl.load_workbook(self.file_path, data_only=True)
                ws = wb.active
                rows = list(ws.iter_rows(values_only=True))
                if not rows or len(rows) < 2:
                    self.finished.emit(False, {}, "El archivo seleccionado no contiene filas de datos.")
                    return

                # Detección inteligente de la fila de encabezados (revisar primeras 10 filas)
                header_row_idx = 0
                idx_int = -1
                idx_desc = -1
                raw_headers = []

                for r_idx in range(min(10, len(rows))):
                    candidate_headers = [str(h or "").strip().upper() for h in rows[r_idx]]
                    def test_idx(*names):
                        for name in names:
                            for c_i, ch in enumerate(candidate_headers):
                                if ch == name.upper() or ch.startswith(name.upper()):
                                    return c_i
                        return -1

                    t_int = test_idx("INTERIOR", "VIVIENDA", "DEPTO", "DEPARTAMENTO", "CASA", "UPE", "INMUEBLE")
                    t_desc = test_idx("DESCRIPCION", "DESCRIPCIÓN", "NOTAS", "OBSERVACIONES", "CONCEPTO", "DETALLE")
                    if t_int != -1 or t_desc != -1:
                        header_row_idx = r_idx
                        raw_headers = candidate_headers
                        break

                if not raw_headers:
                    header_row_idx = 0
                    raw_headers = [str(h or "").strip().upper() for h in rows[0]]

                def get_col_idx(*names):
                    for name in names:
                        for c_i, ch in enumerate(raw_headers):
                            if ch == name.upper() or ch.startswith(name.upper()):
                                return c_i
                    return -1

                idx_int = get_col_idx("INTERIOR", "VIVIENDA", "DEPTO", "DEPARTAMENTO", "CASA", "UPE", "INMUEBLE")
                idx_ed = get_col_idx("EDIFICIO", "EDIF", "TORRE", "EDIF.")
                idx_niv = get_col_idx("NIVEL", "PISO", "PLANTA")
                idx_nof = get_col_idx("NO. OFICIAL", "NO._OFICIAL", "NO_OFICIAL", "NUMERO_OFICIAL", "NUMERO OFICIAL", "NUMERO_UPE", "NUM_UPE", "NO.OFICIAL")
                idx_calle = get_col_idx("CALLE", "VIALIDAD", "DIRECCION")
                idx_dest = get_col_idx("DESTINO", "USO")
                idx_esp = get_col_idx("ESPACIO", "TIPO_ESPACIO")
                idx_rpp = get_col_idx("FOLIO", "FOLIO_RPP", "FOLIO_ELECTRONICO", "RPP", "FOLIO REAL", "FOLIO_REAL")
                idx_cat = get_col_idx("CEDULA CATASTRAL", "CEDULA_CATASTRAL", "CEDULA", "CLAVE_CATASTRAL", "CLAVE CATASTRAL", "CATASTRO", "CLAVE_CAT", "CLAVE")
                idx_pad = get_col_idx("PADRON", "CUENTA_PREDIAL", "PREDIAL")
                idx_fec = get_col_idx("FECHA", "FECHA_ALTA", "FECHA_ALTA_FOLIO", "FECHA_ALTA_CC", "FECHA ALTA", "FECHA EXPEDICION", "FECHA_EXPEDICION")
                idx_desc = get_col_idx("DESCRIPCION", "DESCRIPCIÓN", "NOTAS", "OBSERVACIONES", "CONCEPTO", "DETALLE")

                if idx_int == -1 and idx_desc == -1:
                    self.finished.emit(
                        False, {},
                        "El archivo debe contener al menos la columna 'INTERIOR', 'VIVIENDA', 'UPE' o 'DESCRIPCION'."
                    )
                    return

                data_rows = rows[header_row_idx + 1:]
                total_data_rows = max(len(data_rows), 1)

                for r_count, row in enumerate(data_rows):
                    if r_count % 10 == 0 or r_count == total_data_rows - 1:
                        self.progress_updated.emit(
                            r_count + 1,
                            total_data_rows,
                            f"Analizando fila {r_count + 1} de {total_data_rows}..."
                        )

                    if not any(row):
                        continue

                    desc = str(row[idx_desc] or "").strip() if idx_desc != -1 and idx_desc < len(row) and row[idx_desc] is not None else ""
                    interior = str(row[idx_int] or "").strip() if idx_int != -1 and idx_int < len(row) and row[idx_int] is not None else ""
                    nof = str(row[idx_nof] or "").strip() if idx_nof != -1 and idx_nof < len(row) and row[idx_nof] is not None else ""
                    ed = str(row[idx_ed] or "").strip() if idx_ed != -1 and idx_ed < len(row) and row[idx_ed] is not None else ""

                    # Extracción inteligente desde columna DESCRIPCION (ej: "SM-258 MZ-502 L-01 E-A UPE 101 NO OFICIAL 1 PB...")
                    if desc:
                        if not interior:
                            upe_m = re.search(r"\bUPE\s*(\d+)\b", desc, re.IGNORECASE)
                            if upe_m:
                                interior = upe_m.group(1).strip()
                        if not nof:
                            nof_m = re.search(r"\bNO\.?\s*OFICIAL\s*(\d+)\b", desc, re.IGNORECASE)
                            if nof_m:
                                nof = nof_m.group(1).strip()
                        if not ed:
                            ed_m = re.search(r"\bE-([A-Z0-9]+)\b", desc, re.IGNORECASE)
                            if ed_m:
                                ed = ed_m.group(1).strip().upper()

                    # Si después de analizar no tenemos interior pero sí número oficial, usar número oficial como interior
                    if not interior and nof:
                        interior = nof

                    if not interior:
                        continue

                    niv = str(row[idx_niv] or "").strip() if idx_niv != -1 and idx_niv < len(row) and row[idx_niv] is not None else ""
                    calle = str(row[idx_calle] or "").strip() if idx_calle != -1 and idx_calle < len(row) and row[idx_calle] is not None else ""
                    dest = str(row[idx_dest] or "HABITACIONAL").strip().upper() if idx_dest != -1 and idx_dest < len(row) and row[idx_dest] is not None else "HABITACIONAL"
                    esp_raw = str(row[idx_esp] or "").strip().upper() if idx_esp != -1 and idx_esp < len(row) and row[idx_esp] is not None else ""
                    if dest == "SERVICIO":
                        esp = esp_raw if esp_raw and esp_raw not in ["VIVIENDA", ""] else "AREAS_COMUNES"
                    else:
                        esp = esp_raw or "VIVIENDA"
                    rpp = str(row[idx_rpp] or "").strip() if idx_rpp != -1 and idx_rpp < len(row) and row[idx_rpp] is not None else ""
                    cat = str(row[idx_cat] or "").strip() if idx_cat != -1 and idx_cat < len(row) and row[idx_cat] is not None else ""
                    pad = str(row[idx_pad] or "").strip() if idx_pad != -1 and idx_pad < len(row) and row[idx_pad] is not None else ""
                    raw_fec = row[idx_fec] if idx_fec != -1 and row[idx_fec] is not None else None
                    fec_str = ""
                    if raw_fec is not None:
                        if hasattr(raw_fec, "strftime"):
                            fec_str = raw_fec.strftime("%Y-%m-%d")
                        else:
                            fec_str = str(raw_fec).strip()
                    desc = str(row[idx_desc] or "").strip() if idx_desc != -1 and row[idx_desc] is not None else ""

                    rec_item = {
                        "interior": interior,
                        "edificio": ed,
                        "nivel": niv,
                        "numero_oficial": nof,
                        "calle": calle,
                        "destino": dest,
                        "espacio": esp,
                        "folio_electronico": rpp,
                        "clave_catastral": cat,
                        "padron": pad,
                        "fecha": fec_str,
                        "descripcion": desc
                    }
                    if self.modo == "RPP" and fec_str:
                        rec_item["fecha_alta_folio"] = fec_str
                    elif self.modo == "CATASTRO" and fec_str:
                        rec_item["fecha_alta_cc"] = fec_str

                    records.append(rec_item)

            if not records:
                self.finished.emit(False, {}, "No se detectaron registros con Interior válido.")
                return

            # Detección exhaustiva de duplicados internos dentro del mismo archivo
            seen_file_keys = {}
            seen_folios = {}
            seen_claves = {}
            file_duplicates = []
            duplicate_folios = []
            duplicate_claves = []

            for idx_r, r in enumerate(records):
                # 1. Duplicidad física de unidad: (edificio, interior)
                k_file = ((r.get("edificio") or "").strip().upper(), (r.get("interior") or "").strip().upper())
                if k_file in seen_file_keys:
                    prev_idx = seen_file_keys[k_file]
                    file_duplicates.append((f"{k_file[0]} Int {k_file[1]}", prev_idx, idx_r))
                else:
                    seen_file_keys[k_file] = idx_r

                # 2. Duplicidad de Folio Electrónico RPP dentro del archivo
                fol_raw = str(r.get("folio_electronico") or "").strip().upper()
                if fol_raw and fol_raw not in ("-", "N/A", "NONE"):
                    if fol_raw in seen_folios:
                        prev_idx = seen_folios[fol_raw]
                        duplicate_folios.append((fol_raw, prev_idx, idx_r))
                    else:
                        seen_folios[fol_raw] = idx_r

                # 3. Duplicidad de Clave Catastral dentro del archivo
                cc_raw = str(r.get("clave_catastral") or "").strip().upper()
                if cc_raw and cc_raw not in ("-", "N/A", "NONE"):
                    if cc_raw in seen_claves:
                        prev_idx = seen_claves[cc_raw]
                        duplicate_claves.append((cc_raw, prev_idx, idx_r))
                    else:
                        seen_claves[cc_raw] = idx_r

            # En Carga Inicial: si hay unidades duplicadas internamente, se bloquea la operación para no corromper la BD
            if file_duplicates and self.modo == "INICIAL":
                ejemplos = ", ".join([d[0] for d in file_duplicates[:5]])
                self.finished.emit(
                    False, {},
                    f"El archivo contiene {len(file_duplicates)} registros duplicados internamente:\n"
                    f"({ejemplos}{'...' if len(file_duplicates) > 5 else ''}).\n\n"
                    "Por favor elimine las filas duplicadas en el archivo antes de continuar."
                )
                return

            # Validación de UPEs autorizadas vs total en archivo y capacidad restante
            tot_upes = 0
            if self.regimen_id:
                if self.api_client.connect_via_api:
                    all_regs = self.api_client.request("GET", "/api/admin/data/satc_regimenes")
                    target = next((r for r in all_regs if r.get("regimen_id") == self.regimen_id), None)
                    tot_upes = target.get("total_upes", 0) if target else 0
                else:
                    with self.db_connector.get_session() as session:
                        svc = HousingLifeCycleService(session)
                        target = svc.regimen_repo.get_by_id(self.regimen_id)
                        tot_upes = target.total_upes if target else 0

                # En Carga Inicial es donde se restringe estrictamente la capacidad del macro-lote
                if self.modo == "INICIAL" and len(records) > tot_upes:
                    self.finished.emit(
                        False, {},
                        f"El archivo contiene {len(records)} unidades, lo cual SUPERA la capacidad autorizada "
                        f"del Régimen ({tot_upes} UPEs).\n\n"
                        f"Por favor verifique el archivo para no exceder las {tot_upes} UPEs del macro-lote."
                    )
                    return

            # Consultar existentes en BD para marcar si es NUEVO, VINCULABLE o DUPLICADO
            if self.api_client.connect_via_api:
                params = {"desarrollo_id": self.desarrollo_id}
                if self.regimen_id:
                    params["regimen_id"] = self.regimen_id
                v_existentes = self.api_client.request("GET", f"/api/admin/satc/viviendas/{self.desarrollo_id}", data=params)
            else:
                with self.db_connector.get_session() as session:
                    svc = HousingLifeCycleService(session)
                    if self.regimen_id:
                        v_existentes = svc.get_viviendas_by_regimen(self.regimen_id)
                    else:
                        v_existentes = svc.get_viviendas_by_desarrollo(self.desarrollo_id)

            # Normalizar existentes:
            # - Tupla conjunta completa: (edificio, interior, no_oficial)
            # - Tupla física: (edificio, interior)
            # - Mapa de servicio: (interior, no_oficial) para unidades SERVICIO o sin edificio
            # - Mapa por folio electrónico RPP
            # - Conteo por interior y número oficial
            tupla_completa_map = {}
            servicio_map = {}
            existentes_map = {}
            interior_map = {}
            interior_counts = {}
            nof_map = {}
            nof_counts = {}
            rpp_map = {}
            for v in v_existentes:
                dest_v = str(v.get("destino") or "").strip().upper()
                ed_raw = (v.get("edificio") or "").strip().upper()
                ed_clean = "" if ed_raw in ("-", "N/A", "NONE", "") else re.sub(r'^(EDIFICIO|TORRE)\s*', '', ed_raw).strip()
                int_clean = (v.get("interior") or "").strip().upper()
                nof_raw = str(v.get("numero_oficial") or "").strip().upper()
                digits_match = re.search(r'\b(\d+)\b', nof_raw)
                nof_clean = digits_match.group(1).lstrip("0") or "0" if digits_match else nof_raw

                # Indexar tupla completa conjunta (Edificio, UPE/Interior, Número Oficial)
                if int_clean:
                    tupla_completa_map[(ed_clean, int_clean, nof_clean)] = v
                    tupla_completa_map[(ed_raw if ed_raw != "-" else "", int_clean, nof_clean)] = v
                    if dest_v == "SERVICIO" or not ed_clean:
                        servicio_map[(int_clean, nof_clean)] = v
                        servicio_map[(int_clean, nof_raw)] = v

                existentes_map[(ed_clean, int_clean)] = v
                if ed_raw and ed_raw != "-":
                    existentes_map[(ed_raw, int_clean)] = v
                interior_counts[int_clean] = interior_counts.get(int_clean, 0) + 1
                interior_map[int_clean] = v

                # Indexar por Folio Electrónico RPP (Prioridad Primaria Soberana)
                fol_raw = str(v.get("folio_electronico") or "").strip().upper()
                if fol_raw and fol_raw not in ("-", "N/A", "NONE"):
                    rpp_map[fol_raw] = v

                # Indexar por Número Oficial
                if nof_raw and nof_raw not in ("-", "N/A", "NONE"):
                    nof_map[nof_raw] = v
                    nof_counts[nof_raw] = nof_counts.get(nof_raw, 0) + 1
                    if nof_clean != nof_raw:
                        if nof_clean not in nof_map:
                            nof_map[nof_clean] = v
                        nof_counts[nof_clean] = nof_counts.get(nof_clean, 0) + 1

            # Clasificación de registros y estados
            nuevos_cnt = 0
            duplicados_cnt = 0
            processed_records = []
            matched_v_ids = set()

            for r in records:
                raw_ed = (r.get("edificio") or "").strip().upper()
                ed_clean = "" if raw_ed in ("-", "N/A", "NONE", "") else re.sub(r'^(EDIFICIO|TORRE)\s*', '', raw_ed).strip()
                int_clean = (r.get("interior") or "").strip().upper()
                raw_nof = str(r.get("numero_oficial") or "").strip().upper()
                digits_nof = re.search(r'\b(\d+)\b', raw_nof)
                nof_clean = digits_nof.group(1).lstrip("0") or "0" if digits_nof else raw_nof
                raw_folio = str(r.get("folio_electronico") or "").strip().upper()

                target_v = None

                # 1. VALIDACIÓN PRIMARIA: Coincidencia por Folio Electrónico RPP fidedigno
                if raw_folio and raw_folio in rpp_map:
                    target_v = rpp_map.get(raw_folio)

                # 2. VALIDACIÓN POR NÚMERO OFICIAL UNÍVOCO EN EL RÉGIMEN (Prioridad Primaria en RPP/Catastro)
                if not target_v and raw_nof:
                    if nof_counts.get(raw_nof) == 1:
                        target_v = nof_map.get(raw_nof)
                    elif nof_clean and nof_counts.get(nof_clean) == 1:
                        target_v = nof_map.get(nof_clean)

                # 3. VALIDACIÓN PARA UNIDADES DE SERVICIO / SIN EDIFICIO: (Interior, Número Oficial)
                if not target_v and int_clean and raw_nof and (not ed_clean or raw_ed in ("-", "N/A", "NONE", "")):
                    target_v = servicio_map.get((int_clean, nof_clean)) or servicio_map.get((int_clean, raw_nof))

                # 4. VALIDACIÓN SECUNDARIA: Coincidencia conjunta por (Edificio, UPE/Interior, Número Oficial)
                if not target_v and int_clean and nof_clean:
                    target_v = tupla_completa_map.get((ed_clean, int_clean, nof_clean)) or tupla_completa_map.get((raw_ed, int_clean, nof_clean))

                # 5. VALIDACIÓN FÍSICA DIRECTA: Coincidencia por combinación de (Edificio, UPE/Interior)
                if not target_v and (ed_clean or raw_ed):
                    target_v = existentes_map.get((ed_clean, int_clean)) or existentes_map.get((raw_ed, int_clean))

                # 6. VALIDACIÓN CRUZADA: Coincidencia con campos transpuestos (Interior <-> Edificio)
                if not target_v and (ed_clean or raw_ed) and int_clean:
                    target_v = existentes_map.get((int_clean, ed_clean)) or existentes_map.get((int_clean, raw_ed))

                # 7. Respaldo por (interior, número oficial) general
                if not target_v and int_clean and raw_nof:
                    target_v = servicio_map.get((int_clean, nof_clean)) or servicio_map.get((int_clean, raw_nof))

                # 8. Respaldo para áreas comunes sin edificio asignado con interior unívoco
                if not target_v and not ed_clean and not raw_ed and interior_counts.get(int_clean) == 1:
                    target_v = interior_map.get(int_clean)

                # 9. Inferencia Catastral: Si es CATASTRO y no hubo coincidencia, deducir por sufijo de Clave Catastral
                if not target_v and self.modo == "CATASTRO":
                    raw_cc = str(r.get("clave_catastral") or "").strip()
                    if raw_cc.startswith("60") and len(raw_cc) == 18:
                        cc_nof = raw_cc[-4:].lstrip("0") or "0"
                        if nof_counts.get(cc_nof) == 1:
                            target_v = nof_map.get(cc_nof)

                if self.modo in ("RPP", "CATASTRO"):
                    if target_v:
                        status_str = "VINCULABLE"
                        nuevos_cnt += 1
                        v_id = target_v.get("vivienda_id")
                        if v_id is not None:
                            matched_v_ids.add(v_id)
                        # Preservar atributos fidedignos de la vivienda en BD (Destino, Espacio, Edificio real y Folio RPP)
                        r_display = dict(r)
                        r_display["destino"] = target_v.get("destino") or r.get("destino") or "HABITACIONAL"
                        r_display["espacio"] = target_v.get("espacio") or r.get("espacio") or "VIVIENDA"
                        r_display["edificio"] = target_v.get("edificio") or ""
                        if r_display["edificio"] == "-":
                            r_display["edificio"] = ""
                        r_display["nivel"] = target_v.get("nivel") or r.get("nivel") or ""
                        r_display["numero_oficial"] = target_v.get("numero_oficial") or r.get("numero_oficial") or ""
                        r_display["calle"] = target_v.get("calle") or r.get("calle") or ""
                        r_display["descripcion"] = target_v.get("descripcion") or r.get("descripcion") or ""

                        # Blindaje ante sellos: Si el oficio traía el folio RPP obstaculizado por sello, rescatar el fidedigno de la BD
                        if not r_display.get("folio_electronico") and target_v.get("folio_electronico"):
                            r_display["folio_electronico"] = target_v.get("folio_electronico")
                            r["folio_electronico"] = target_v.get("folio_electronico")

                        # Validación Metadata-Driven de Clave Catastral
                        if self.modo == "CATASTRO":
                            cc_val = str(r.get("clave_catastral") or "").strip()
                            if cc_val:
                                from satc.src.services.metadata_extraction_service import MetadataExtractionService
                                # Resolver sesión si está disponible o validar con motor desacoplado
                                with self.db_connector.get_session() as s_meta:
                                    meta_svc = MetadataExtractionService(s_meta)
                                    # Por defecto validar Benito Juárez (ID: 2) si no se especifica otro municipio
                                    es_val, cc_clean, err_msg = meta_svc.es_clave_catastral_valida(cc_val, municipio_id=2)
                                    if not es_val:
                                        # Si no cumple Cancún, validar si cumple Playa (ID: 8) o Tulum (ID: 9)
                                        es_val_p, _, _ = meta_svc.es_clave_catastral_valida(cc_val, municipio_id=8)
                                        es_val_t, _, _ = meta_svc.es_clave_catastral_valida(cc_val, municipio_id=9)
                                        if not es_val_p and not es_val_t:
                                            r_display["descripcion"] = f"⚠️ [CC INVÁLIDA: {err_msg}] {r_display['descripcion']}".strip()
                    else:
                        status_str = "NO ENCONTRADO"
                        duplicados_cnt += 1
                        r_display = r
                else:
                    if target_v:
                        status_str = "DUPLICADO"
                        duplicados_cnt += 1
                    else:
                        status_str = "NUEVO"
                        nuevos_cnt += 1
                    r_display = r

                # Detección de colisiones de duplicados internos en este registro específico
                advertencias_dup = []
                if self.modo == "RPP" and raw_folio:
                    if any(df[0] == raw_folio for df in duplicate_folios):
                        advertencias_dup.append(f"⚠️ [FOLIO DUPLICADO EN ARCHIVO: {raw_folio}]")
                elif self.modo == "CATASTRO":
                    cc_raw_val = str(r.get("clave_catastral") or "").strip().upper()
                    if cc_raw_val and any(dc[0] == cc_raw_val for dc in duplicate_claves):
                        advertencias_dup.append(f"⚠️ [CC DUPLICADA EN ARCHIVO: {cc_raw_val}]")
                    if raw_folio and any(df[0] == raw_folio for df in duplicate_folios):
                        advertencias_dup.append(f"⚠️ [FOLIO DUPLICADO EN ARCHIVO: {raw_folio}]")

                if advertencias_dup:
                    desc_orig = r_display.get("descripcion") or ""
                    r_display["descripcion"] = f"{' '.join(advertencias_dup)} {desc_orig}".strip()

                processed_records.append({
                    "data": r_display,
                    "status": status_str
                })

            # Detectar viviendas de la base de datos que NO vinieron en el archivo cargado
            faltantes_bd = []
            if self.modo in ("RPP", "CATASTRO") and v_existentes:
                for v in v_existentes:
                    v_id = v.get("vivienda_id")
                    if v_id is not None and v_id not in matched_v_ids:
                        faltantes_bd.append(v)

            elapsed_sec = round(time.time() - start_time, 2)
            result_data = {
                "records": records,
                "processed_records": processed_records,
                "faltantes_bd": faltantes_bd,
                "file_name": os.path.basename(self.file_path),
                "nuevos_cnt": nuevos_cnt,
                "duplicados_cnt": duplicados_cnt,
                "tot_upes": tot_upes,
                "duplicate_folios": duplicate_folios,
                "duplicate_claves": duplicate_claves,
                "file_duplicates": file_duplicates,
                "elapsed_time": elapsed_sec
            }
            self.finished.emit(True, result_data, "")
        except Exception as e:
            self.finished.emit(False, {}, str(e))


class ViviendasLoadWorker(QThread):
    """Worker thread para cargar viviendas y capacidad del régimen en segundo plano sin congelar la UI."""
    finished = Signal(bool, dict, str)

    def __init__(self, api_client, db_connector, desarrollo_id: int, regimen_id: Optional[int]):
        super().__init__()
        self.api_client = api_client
        self.db_connector = db_connector
        self.desarrollo_id = desarrollo_id
        self.regimen_id = regimen_id

    def run(self):
        try:
            tot_u = 0
            # 1. Consultar capacidad del régimen si aplica
            if self.regimen_id:
                if self.api_client.connect_via_api:
                    all_regs = self.api_client.request("GET", "/api/admin/data/satc_regimenes")
                    target = next((r for r in all_regs if r.get("regimen_id") == self.regimen_id), None)
                    tot_u = target.get("total_upes", 0) if target else 0
                else:
                    with self.db_connector.get_session() as session:
                        svc = HousingLifeCycleService(session)
                        target = svc.regimen_repo.get_by_id(self.regimen_id)
                        tot_u = target.total_upes if target else 0

            # 2. Consultar viviendas registradas
            if self.api_client.connect_via_api:
                params = {"desarrollo_id": self.desarrollo_id}
                if self.regimen_id:
                    params["regimen_id"] = self.regimen_id
                viviendas = self.api_client.request("GET", f"/api/admin/satc/viviendas/{self.desarrollo_id}", data=params)
            else:
                with self.db_connector.get_session() as session:
                    svc = HousingLifeCycleService(session)
                    if self.regimen_id:
                        viviendas = svc.get_viviendas_by_regimen(self.regimen_id)
                    else:
                        viviendas = svc.get_viviendas_by_desarrollo(self.desarrollo_id)

            res = {
                "viviendas": viviendas,
                "tot_u": tot_u,
                "regimen_id": self.regimen_id,
                "desarrollo_id": self.desarrollo_id
            }
            self.finished.emit(True, res, "")
        except Exception as e:
            self.finished.emit(False, {}, str(e))


class ViviendasImportWorker(QThread):
    """Worker thread para ejecutar la importación y vinculación progresiva sin congelar la interfaz."""
    finished = Signal(bool, dict, str)

    def __init__(self, api_client, db_connector, desarrollo_id: int, regimen_id: Optional[int], records: list, usuario_id: int, sesion_id: Optional[int] = None, modo: str = "INICIAL"):
        super().__init__()
        self.api_client = api_client
        self.db_connector = db_connector
        self.desarrollo_id = desarrollo_id
        self.regimen_id = regimen_id
        self.records = records
        self.usuario_id = usuario_id
        self.sesion_id = sesion_id
        self.modo = modo  # "INICIAL", "RPP", "CATASTRO"

    def run(self):
        import time
        start_time = time.time()
        try:
            if self.api_client.connect_via_api:
                endpoint = "/api/admin/satc/viviendas/bulk-import"
                if self.modo == "RPP":
                    endpoint = f"/api/admin/satc/viviendas/{self.regimen_id}/vincular-rpp"
                elif self.modo == "CATASTRO":
                    endpoint = f"/api/admin/satc/viviendas/{self.regimen_id}/vincular-catastro"

                payload = {
                    "desarrollo_id": self.desarrollo_id,
                    "regimen_id": self.regimen_id,
                    "usuario_id": self.usuario_id,
                    "sesion_id": self.sesion_id,
                    "records": self.records
                }
                res = self.api_client.request("POST", endpoint, json=payload)
            else:
                with self.db_connector.get_session() as session:
                    svc = HousingLifeCycleService(session)
                    if self.modo == "RPP":
                        res = svc.vincular_folios_rpp_excel(
                            regimen_id=self.regimen_id,
                            records=self.records,
                            usuario_id=self.usuario_id,
                            sesion_id=self.sesion_id
                        )
                    elif self.modo == "CATASTRO":
                        res = svc.vincular_catastro_excel(
                            regimen_id=self.regimen_id,
                            records=self.records,
                            usuario_id=self.usuario_id,
                            sesion_id=self.sesion_id
                        )
                    else:
                        res = svc.importar_viviendas_excel(
                            desarrollo_id=self.desarrollo_id,
                            records=self.records,
                            regimen_id=self.regimen_id,
                            usuario_id=self.usuario_id,
                            sesion_id=self.sesion_id
                        )
            elapsed_sec = round(time.time() - start_time, 2)
            if isinstance(res, dict):
                res["elapsed_time"] = elapsed_sec
            self.finished.emit(True, res, "")
        except Exception as e:
            self.finished.emit(False, {}, str(e))


class SatcCatalogsView(QWidget):
    """Sub-vista integrada para la administración de catálogos maestros de SATC."""
    action_state_changed = Signal()

    def __init__(self, db_connector, current_user_id, current_sesion_id, can_edit, parent=None):
        super().__init__(parent)
        self.db_connector = db_connector
        self.current_user_id = current_user_id
        self.current_sesion_id = current_sesion_id
        self.can_edit = can_edit

        from sar.src.storage.api_client import APIClient
        self.api_client = APIClient()

        self.layout = QHBoxLayout(self)
        self.layout.setContentsMargins(0, 0, 0, 0)

        self.desarrollos_list = []
        self._v_loading_dialog: Optional[GLLoadingDialog] = None
        self._build_ui()

    def _build_ui(self):
        self.tabs = QTabWidget(self)
        self.tabs.setObjectName("seamlessTabWidget")
        self.tabs.tabBar().setVisible(False)
        self.layout.addWidget(self.tabs)

        # ---------------------------------------------------------
        # Pestaña 1: Tipos de Trámite Municipal
        # ---------------------------------------------------------
        self.tab_tipos = QWidget()
        lay_tipos = QHBoxLayout(self.tab_tipos)
        lay_tipos.setContentsMargins(0, 0, 0, 0)
        self.tbl_tipos = CrudTablePanel("Catálogo de Tipos de Trámite SATC")
        self.tbl_tipos.setup_table(
            ["ID", "Código", "Alias", "Nombre", "Aplica A", "Etapa Disparador", "Vigencia (Días)", "Obligatorio", "Estado"],
            ["tipo_tramite_id", "codigo", "alias", "nombre", "aplica_a", "etapa_disparador", "dias_vigencia", "es_obligatorio", "activo"]
        )
        self.tbl_tipos.add_requested.connect(self._on_new_tipo)
        self.tbl_tipos.edit_requested.connect(self._on_edit_tipo)
        self.tbl_tipos.btn_add.setVisible(self.can_edit)
        self.tbl_tipos.btn_edit.setVisible(self.can_edit)
        lay_tipos.addWidget(self.tbl_tipos)
        self.tabs.addTab(self.tab_tipos, "Tipos de Trámite")

        # ---------------------------------------------------------
        # Pestaña 2: Estados de Trámite
        # ---------------------------------------------------------
        self.tab_est_tramite = QWidget()
        lay_est_tramite = QHBoxLayout(self.tab_est_tramite)
        lay_est_tramite.setContentsMargins(0, 0, 0, 0)
        self.tbl_est_tramite = CrudTablePanel("Estados de Trámite SATC")
        self.tbl_est_tramite.setup_table(
            ["ID", "Código", "Nombre", "Fase", "Bloqueante", "Terminal", "Estado"],
            ["estado_tramite_id", "codigo", "nombre", "fase", "es_bloqueante", "es_terminal", "activo"]
        )
        self.tbl_est_tramite.add_requested.connect(self._on_new_est_tramite)
        self.tbl_est_tramite.edit_requested.connect(self._on_edit_est_tramite)
        self.tbl_est_tramite.btn_add.setVisible(self.can_edit)
        self.tbl_est_tramite.btn_edit.setVisible(self.can_edit)
        lay_est_tramite.addWidget(self.tbl_est_tramite)
        self.tabs.addTab(self.tab_est_tramite, "Estados de Trámite")

        # ---------------------------------------------------------
        # Pestaña 3: Estados de Vivienda (Ciclo Físico-Comercial)
        # ---------------------------------------------------------
        self.tab_est_vivienda = QWidget()
        lay_est_vivienda = QHBoxLayout(self.tab_est_vivienda)
        lay_est_vivienda.setContentsMargins(0, 0, 0, 0)
        self.tbl_est_vivienda = CrudTablePanel("Estados de Vivienda / UPE")
        self.tbl_est_vivienda.setup_table(
            ["ID", "Código", "Nombre", "Descripción", "Orden", "Permite Venta", "Permite Escrituración", "Estado"],
            ["estado_vivienda_id", "codigo", "nombre", "descripcion", "orden_flujo", "permite_venta", "permite_escrituracion", "activo"]
        )
        self.tbl_est_vivienda.add_requested.connect(self._on_new_est_vivienda)
        self.tbl_est_vivienda.edit_requested.connect(self._on_edit_est_vivienda)
        self.tbl_est_vivienda.btn_add.setVisible(self.can_edit)
        self.tbl_est_vivienda.btn_edit.setVisible(self.can_edit)
        lay_est_vivienda.addWidget(self.tbl_est_vivienda)
        self.tabs.addTab(self.tab_est_vivienda, "Estados de Vivienda")

        # ---------------------------------------------------------
        # Pestaña 4: Regímenes mz
        # ---------------------------------------------------------
        self.tab_regimenes = QWidget()
        lay_regimenes = QHBoxLayout(self.tab_regimenes)
        lay_regimenes.setContentsMargins(0, 0, 0, 0)
        self.tbl_regimenes = CrudTablePanel("Regímenes mz")
        self.tbl_regimenes.setup_table(
            ["ID", "Desarrollo", "Nombre Régimen", "SM", "MZ", "LT", "Total UPEs", "Generadas", "Licencia Matriz", "Clave Catastral", "Estado"],
            ["regimen_id", "desarrollo_nombre", "nombre_regimen", "sm", "mz", "lt", "total_upes", "upes_generadas", "licencia_construccion_matriz", "clave_catastral_matriz", "activo"]
        )
        self.tbl_regimenes.add_requested.connect(self._on_new_regimen)
        self.tbl_regimenes.edit_requested.connect(self._on_edit_regimen)
        self.tbl_regimenes.btn_add.setVisible(self.can_edit)
        self.tbl_regimenes.btn_edit.setVisible(self.can_edit)
        lay_regimenes.addWidget(self.tbl_regimenes)
        self.tabs.addTab(self.tab_regimenes, "Regímenes mz")

        # ---------------------------------------------------------
        # Pestaña 5: Importación y Padrón de Viviendas
        # ---------------------------------------------------------
        self.tab_import_viviendas = QWidget()
        lay_import_viv = QVBoxLayout(self.tab_import_viviendas)
        lay_import_viv.setContentsMargins(12, 12, 12, 12)
        lay_import_viv.setSpacing(12)

        # 1. Barra de filtros y controles (Elástica y adaptativa a pantallas pequeñas)
        ctl_box = QHBoxLayout()
        ctl_box.setSpacing(10)

        ctl_box.addWidget(CustomLabel("Desarrollo *:", variant="body"))
        self.cmb_v_desarrollo = CustomComboBox()
        self.cmb_v_desarrollo.setMinimumWidth(160)
        self.cmb_v_desarrollo.setSizePolicy(QSizePolicy.Expanding, QSizePolicy.Fixed)
        self.cmb_v_desarrollo.currentIndexChanged.connect(self._on_v_desarrollo_changed)
        ctl_box.addWidget(self.cmb_v_desarrollo, 1)

        ctl_box.addWidget(CustomLabel("Régimen *:", variant="body"))
        self.cmb_v_regimen = CustomComboBox()
        self.cmb_v_regimen.setMinimumWidth(180)
        self.cmb_v_regimen.setSizePolicy(QSizePolicy.Expanding, QSizePolicy.Fixed)
        self.cmb_v_regimen.currentIndexChanged.connect(self._on_v_regimen_changed)
        ctl_box.addWidget(self.cmb_v_regimen, 1)

        ctl_box.addWidget(CustomLabel("Operación:", variant="body"))
        self.cmb_v_modo = CustomComboBox()
        self.cmb_v_modo.setMinimumWidth(180)
        self.cmb_v_modo.setSizePolicy(QSizePolicy.Preferred, QSizePolicy.Fixed)
        self.cmb_v_modo.addItem("-- Selecciona Operación --", None)
        self.cmb_v_modo.addItem("1. Carga Inicial (Física)", "INICIAL")
        self.cmb_v_modo.addItem("2. Vincular Folios RPP", "RPP")
        self.cmb_v_modo.addItem("3. Vincular CC / Catastro", "CATASTRO")
        self.cmb_v_modo.currentIndexChanged.connect(self._on_v_modo_changed)
        ctl_box.addWidget(self.cmb_v_modo)

        ctl_box.addStretch()

        # Botones operativos desacoplados (sus acciones se invocan desde la Toolbar superior)
        self.btn_v_plantilla = CustomButton("Plantilla Excel", is_secondary=True, icon_name="excel")
        self.btn_v_plantilla.clicked.connect(self._on_v_descargar_plantilla)
        self.btn_v_plantilla.hide()

        self.btn_v_seleccionar = CustomButton("Cargar Archivo", is_secondary=True, icon_name="documento_cargar")
        self.btn_v_seleccionar.clicked.connect(self._on_v_seleccionar_archivo)
        self.btn_v_seleccionar.hide()

        self.btn_v_exportar_datos = CustomButton("Exportar Datos", is_secondary=True, icon_name="excel")
        self.btn_v_exportar_datos.clicked.connect(self._on_v_exportar_datos)
        self.btn_v_exportar_datos.setEnabled(False)
        self.btn_v_exportar_datos.hide()

        self.btn_v_limpiar = CustomButton("Limpiar", is_secondary=True, icon_name="cancelar")
        self.btn_v_limpiar.clicked.connect(self._on_v_limpiar)
        self.btn_v_limpiar.hide()

        self.btn_v_ejecutar = CustomButton("Aplicar Cambios", icon_name="guardar")
        self.btn_v_ejecutar.clicked.connect(self._on_v_ejecutar_importacion)
        self.btn_v_ejecutar.setEnabled(False)
        self.btn_v_ejecutar.setVisible(self.can_edit)
        self.btn_v_ejecutar.hide()

        lay_import_viv.addLayout(ctl_box)

        # 2. Information Row / Metrics Bar (Limpio, transparente, perfectamente distribuido y uniforme)
        self.kpi_v_frame = QFrame()
        self.kpi_v_frame.setStyleSheet("""
            QFrame {
                background-color: transparent;
                border: none;
            }
        """)
        kpi_lay = QHBoxLayout(self.kpi_v_frame)
        kpi_lay.setContentsMargins(0, 4, 0, 4)
        kpi_lay.setSpacing(6)

        self.kpi_containers = {}

        def _make_clickable_kpi(dot_color: str, initial_text: str, kpi_key: str, tooltip_text: str) -> tuple[QFrame, QLabel, CustomLabel]:
            container = QFrame()
            container.setObjectName(f"kpi_{kpi_key}")
            container.setCursor(Qt.PointingHandCursor)
            container.setToolTip(tooltip_text)
            container.setStyleSheet(f"""
                QFrame#{container.objectName()} {{
                    background-color: transparent;
                    border: 1px solid transparent;
                    border-radius: 4px;
                    padding: 2px 4px;
                }}
                QFrame#{container.objectName()}:hover {{
                    background-color: {Colors.NEUTRAL_BG};
                    border: 1px solid {Colors.BORDER_LIGHT};
                }}
            """)

            lay = QHBoxLayout(container)
            lay.setContentsMargins(4, 2, 4, 2)
            lay.setSpacing(6)

            dot = QLabel()
            dot.setFixedSize(10, 10)
            dot.setStyleSheet(f"background-color: {dot_color}; border-radius: 5px; border: none;")

            lbl = CustomLabel(initial_text, variant="body")
            lbl.setStyleSheet("background: transparent;")

            lay.addWidget(dot)
            lay.addWidget(lbl)

            # Conectar clic de ratón al despachador de filtro
            container.mousePressEvent = lambda e, k=kpi_key: self._on_kpi_clicked(k)

            self.kpi_containers[kpi_key] = {
                "container": container,
                "dot": dot,
                "lbl": lbl,
                "dot_color": dot_color
            }
            return container, dot, lbl

        # Puntos e indicadores KPI interactivos (Atomic Design: Tokens de Colors)
        self.kpi_total_frame, self.dot_vkpi_total, self.lbl_vkpi_total_arch = _make_clickable_kpi(
            Colors.TEXT_LIGHT_MUTED, "En Archivo: 0", "TOTAL", "Clic para mostrar todos los registros del archivo"
        )
        self.kpi_nuevos_frame, self.dot_vkpi_nuevos, self.lbl_vkpi_nuevos = _make_clickable_kpi(
            Colors.SUCCESS, "Nuevos / Vinculables: 0", "VINCULABLES", "Clic para filtrar solo registros vinculables"
        )
        self.kpi_duplicados_frame, self.dot_vkpi_duplicados, self.lbl_vkpi_duplicados = _make_clickable_kpi(
            Colors.ERROR, "Duplicados / Sin Coincidencia: 0", "RECHAZADOS", "Clic para filtrar registros sin coincidencia o duplicados"
        )
        self.kpi_alertas_frame, self.dot_vkpi_alertas, self.lbl_vkpi_alertas = _make_clickable_kpi(
            Colors.WARNING, "Con Advertencia: 0", "ALERTAS", "Clic para filtrar registros con inconsistencia o alerta en el documento"
        )
        self.kpi_alertas_frame.hide()

        self.kpi_faltantes_frame, self.dot_vkpi_faltantes, self.lbl_vkpi_faltantes = _make_clickable_kpi(
            Colors.WARNING, "Faltantes: 0", "FALTANTES", "Clic para filtrar registros de BD que no vinieron en el archivo"
        )
        self.kpi_faltantes_frame.hide()

        # KPIs informativos del régimen y BD (foco / no interactivos de estado)
        def _make_metric_dot(color_val: str) -> QLabel:
            dot = QLabel()
            dot.setFixedSize(10, 10)
            dot.setStyleSheet(f"background-color: {color_val}; border-radius: 5px; border: none;")
            return dot

        self.dot_vkpi_regimen = _make_metric_dot(Colors.ACCENT_VIOLET)
        self.lbl_vkpi_regimen_info = CustomLabel("Total UPE: N/A", variant="body")

        self.dot_vkpi_bd = _make_metric_dot(Colors.ACCENT)
        self.lbl_vkpi_bd = CustomLabel("EnBD: 0", variant="body")

        self.dot_vkpi_rpp = _make_metric_dot(Colors.WARNING)
        self.lbl_vkpi_con_rpp = CustomLabel("Con Folio: 0", variant="body")

        self.dot_vkpi_catastro = _make_metric_dot(Colors.SUCCESS)
        self.lbl_vkpi_con_catastro = CustomLabel("Con CC: 0", variant="body")

        # Layout interactivo equitativo y balanceado
        kpi_lay.addWidget(self.kpi_total_frame)
        kpi_lay.addSpacing(12)

        kpi_lay.addWidget(self.kpi_nuevos_frame)
        kpi_lay.addSpacing(12)

        kpi_lay.addWidget(self.kpi_duplicados_frame)
        kpi_lay.addSpacing(12)

        kpi_lay.addWidget(self.kpi_alertas_frame)
        kpi_lay.addSpacing(12)

        kpi_lay.addWidget(self.kpi_faltantes_frame)
        kpi_lay.addSpacing(12)

        kpi_lay.addWidget(self.dot_vkpi_regimen)
        kpi_lay.addWidget(self.lbl_vkpi_regimen_info)
        kpi_lay.addSpacing(16)

        kpi_lay.addWidget(self.dot_vkpi_bd)
        kpi_lay.addWidget(self.lbl_vkpi_bd)
        kpi_lay.addSpacing(16)

        kpi_lay.addWidget(self.dot_vkpi_rpp)
        kpi_lay.addWidget(self.lbl_vkpi_con_rpp)
        kpi_lay.addSpacing(16)

        kpi_lay.addWidget(self.dot_vkpi_catastro)
        kpi_lay.addWidget(self.lbl_vkpi_con_catastro)

        kpi_lay.addStretch()

        lay_import_viv.addWidget(self.kpi_v_frame)

        # 3. Contenedor de Tarjeta Estandarizada y Tabla (Atomic Design: CustomCard)
        self.card_viviendas = CustomCard("Padrón e Importación de Viviendas", parent=self)
        self.card_viviendas.setObjectName("cardFrame")

        # Indicador de archivo cargado adjunto al título de la tabla
        self.lbl_v_sep = CustomLabel("|", variant="muted")
        self.lbl_v_sep.setStyleSheet(f"color: {Colors.TEXT_DARK_SECONDARY}; font-size: 14px; font-weight: normal; margin: 0 4px;")
        self.card_viviendas.add_header_widget(self.lbl_v_sep)

        self.lbl_vkpi_archivo = CustomLabel("Archivo: Ninguno", variant="muted")
        self.lbl_vkpi_archivo.setStyleSheet(f"font-size: 12px; color: {Colors.TEXT_LIGHT_MUTED}; font-weight: normal;")
        self.card_viviendas.add_header_widget(self.lbl_vkpi_archivo)

        # Selector de Filtro de Estado (Vinculables, Sin Coincidencia, Faltantes, Todos)
        self.cmb_v_filtro_estado = CustomComboBox(parent=self)
        self.cmb_v_filtro_estado.setObjectName("viviendasFiltroEstado")
        self.cmb_v_filtro_estado.setMinimumWidth(180)
        self.cmb_v_filtro_estado.setMaximumWidth(230)
        self.cmb_v_filtro_estado.addItem("Mostrar: Todos", "TODOS")
        self.cmb_v_filtro_estado.currentIndexChanged.connect(self._on_filtro_estado_changed)
        self.card_viviendas.add_header_widget(self.cmb_v_filtro_estado)

        # Campo Buscar alineado al encabezado de la tabla (Homologado con GlCrudTable)
        self.search_viviendas_input = CustomInput(control_size="md", parent=self)
        self.search_viviendas_input.setObjectName("viviendasSearchInput")
        self.search_viviendas_input.setPlaceholderText("Buscar en padrón...")
        self.search_viviendas_input.setMinimumWidth(200)
        self.search_viviendas_input.setMaximumWidth(280)
        self.search_viviendas_input.setClearButtonEnabled(True)
        self.search_viviendas_input.addAction(Icons.search(Colors.TEXT_LIGHT_MUTED), QLineEdit.LeadingPosition)
        self.search_viviendas_input.textChanged.connect(self._on_search_viviendas_text_changed)
        self.card_viviendas.add_header_widget(self.search_viviendas_input)

        self.tbl_viviendas = QTableWidget()
        self.tbl_viviendas.setColumnCount(13)
        self.tbl_viviendas.setHorizontalHeaderLabels([
            "Estado Carga", "Régimen Matriz", "No. Oficial", "Interior", "Edificio", "Nivel", 
            "Destino", "Espacio", "Folio RPP", "Clave Catastral", "Padrón", "Calle", "Notas"
        ])
        self.tbl_viviendas.horizontalHeader().setSectionResizeMode(QHeaderView.Interactive)
        self.tbl_viviendas.horizontalHeader().setStretchLastSection(True)
        self.tbl_viviendas.setHorizontalScrollMode(QAbstractItemView.ScrollPerPixel)
        self.tbl_viviendas.setSelectionBehavior(QAbstractItemView.SelectRows)
        self.tbl_viviendas.setEditTriggers(QAbstractItemView.NoEditTriggers)
        self.tbl_viviendas.setShowGrid(False)
        self.tbl_viviendas.verticalHeader().setVisible(False)
        self.tbl_viviendas.verticalHeader().setDefaultSectionSize(26)
        
        # Anchos iniciales confortables y proporcionales
        col_widths = [110, 160, 90, 80, 80, 70, 110, 90, 130, 140, 110, 160, 200]
        for c_idx, width in enumerate(col_widths):
            self.tbl_viviendas.setColumnWidth(c_idx, width)
        
        self.card_viviendas.layout.addWidget(self.tbl_viviendas)
        lay_import_viv.addWidget(self.card_viviendas)

        self.tabs.addTab(self.tab_import_viviendas, "Importación de Viviendas")

        # ---------------------------------------------------------
        # Pestaña 6: Reglas de Extracción y Validación (Metadata-Driven)
        # ---------------------------------------------------------
        self.tab_formatos = QWidget()
        lay_formatos = QHBoxLayout(self.tab_formatos)
        lay_formatos.setContentsMargins(0, 0, 0, 0)
        self.tbl_formatos = CrudTablePanel("Reglas y Formatos de Extracción por Municipio (SATC)")
        self.tbl_formatos.setup_table(
            ["ID", "Municipio", "Campo", "Variante", "Prioridad", "Regex Validación", "Longitud", "Sliding Window", "Estado"],
            ["formato_id", "municipio_nombre", "campo_nombre", "nombre_variante", "prioridad", "regex_validacion", "longitud_min", "permite_sliding_window", "activo"]
        )
        self.tbl_formatos.add_requested.connect(self._on_new_formato)
        self.tbl_formatos.edit_requested.connect(self._on_edit_formato)
        self.tbl_formatos.btn_add.setVisible(self.can_edit)
        self.tbl_formatos.btn_edit.setVisible(self.can_edit)
        lay_formatos.addWidget(self.tbl_formatos)
        self.tabs.addTab(self.tab_formatos, "Reglas de Extracción")

        self._v_parsed_records = []
        self._v_worker = None
        self._v_load_worker = None

    # =========================================================================
    # 1. DIALOGOS & ACCIONES: TIPO DE TRÁMITE
    # =========================================================================
    def _create_tipo_dialog(self, title: str) -> CustomDialog:
        dialog = CustomDialog(title, self)
        dialog.setMinimumWidth(560)

        card = QFrame(dialog)
        card.setObjectName("card_tipo")
        card.setStyleSheet(f"""
            QFrame#card_tipo {{
                background-color: {Colors.SLATE_50};
                border: 1px solid {Colors.SLATE_200};
                border-radius: 8px;
                padding: 10px 14px;
            }}
        """)
        lay = QVBoxLayout(card)
        lay.setContentsMargins(0, 0, 0, 0)
        lay.setSpacing(8)

        lbl = CustomLabel("📄 CONFIGURACIÓN DE TIPO DE TRÁMITE", variant="subheader")
        lbl.setStyleSheet(f"font-size: 11px; font-weight: bold; color: {Colors.TEXT_LIGHT_PRIMARY}; margin-bottom: 2px;")
        lay.addWidget(lbl)

        form = QFormLayout()
        form.setSpacing(8)
        form.setContentsMargins(0, 0, 0, 0)

        self.inp_t_codigo = CustomInput("Código único (ej. LIC_CONST)", parent=card)
        self.inp_t_codigo.setMaxLength(50)
        self.inp_t_codigo.textEdited.connect(lambda t: self.inp_t_codigo.setText(t.upper()))
        self.inp_t_codigo.text = self.inp_t_codigo.text
        self.inp_t_codigo.set_text = self.inp_t_codigo.setText
        self.inp_t_codigo.set_focus = self.inp_t_codigo.setFocus

        self.inp_t_alias = CustomInput("Alias corto (ej. Licencia Construcción)", parent=card)
        self.inp_t_alias.setMaxLength(100)
        self.inp_t_alias.text = self.inp_t_alias.text
        self.inp_t_alias.set_text = self.inp_t_alias.setText

        self.inp_t_nombre = CustomInput("Nombre oficial del trámite municipal", parent=card)
        self.inp_t_nombre.setMaxLength(200)
        self.inp_t_nombre.text = self.inp_t_nombre.text
        self.inp_t_nombre.set_text = self.inp_t_nombre.setText

        self.cmb_t_aplica = CustomComboBox(parent=card)
        self.cmb_t_aplica.add_items(["AMBOS", "REGIMEN_MATRIZ", "VIVIENDA_INDIVIDUAL"])

        self.cmb_t_etapa = CustomComboBox(parent=card)
        self.cmb_t_etapa.add_items(["CONSTRUCCION", "VENTA", "ESCRITURACION"])

        self.inp_t_vigencia = CustomInput("Días de vigencia (0 = Indefinida)", parent=card)
        self.inp_t_vigencia.setMaxLength(6)
        self.inp_t_vigencia.text = self.inp_t_vigencia.text
        self.inp_t_vigencia.set_text = self.inp_t_vigencia.setText

        self.chk_t_obligatorio = CustomCheckBox("Obligatorio (Bloquea fases subsecuentes)", parent=card)
        self.chk_t_obligatorio.setChecked(True)

        self.chk_t_activo = CustomCheckBox("Activo", parent=card)
        self.chk_t_activo.setChecked(True)

        form.addRow("Código *:", self.inp_t_codigo)
        form.addRow("Alias:", self.inp_t_alias)
        form.addRow("Nombre *:", self.inp_t_nombre)
        form.addRow("Aplica A *:", self.cmb_t_aplica)
        form.addRow("Etapa Disparador *:", self.cmb_t_etapa)
        form.addRow("Vigencia (Días):", self.inp_t_vigencia)
        form.addRow("", self.chk_t_obligatorio)
        form.addRow("", self.chk_t_activo)

        lay.addLayout(form)
        dialog.add_widget(card)

        def _val():
            c = self.inp_t_codigo.text().strip()
            n = self.inp_t_nombre.text().strip()
            dialog.btn_save.setEnabled(bool(c and n))

        self.inp_t_codigo.textChanged.connect(_val)
        self.inp_t_nombre.textChanged.connect(_val)
        _val()

        if not self.can_edit:
            dialog.btn_save.setVisible(False)
            self.inp_t_codigo.setReadOnly(True)
            self.inp_t_alias.setReadOnly(True)
            self.inp_t_nombre.setReadOnly(True)
            self.inp_t_vigencia.setReadOnly(True)

        dialog.btn_save.clicked.disconnect()
        dialog.btn_save.clicked.connect(lambda: self._save_tipo(dialog))
        return dialog

    def _on_new_tipo(self):
        self.current_tipo_id = None
        dialog = self._create_tipo_dialog("Nuevo Tipo de Trámite")
        self.inp_t_vigencia.set_text("0")
        self.inp_t_codigo.setReadOnly(False)
        self.inp_t_codigo.set_focus()
        dialog.exec()

    def _on_edit_tipo(self, data: dict):
        self.current_tipo_id = data.get("tipo_tramite_id")
        dialog = self._create_tipo_dialog(f"Editar Tipo de Trámite: {data.get('codigo')}")
        self.inp_t_codigo.set_text(data.get("codigo", "") or "")
        self.inp_t_alias.set_text(data.get("alias", "") or "")
        self.inp_t_nombre.set_text(data.get("nombre", "") or "")
        self.cmb_t_aplica.set_current_text(data.get("aplica_a", "AMBOS"))
        self.cmb_t_etapa.set_current_text(data.get("etapa_disparador", "CONSTRUCCION"))
        self.inp_t_vigencia.set_text(str(data.get("dias_vigencia", 0) or 0))
        self.chk_t_obligatorio.setChecked(bool(data.get("es_obligatorio", True)))
        self.chk_t_activo.setChecked(bool(data.get("activo", True)))
        self.inp_t_codigo.setReadOnly(False)  # Permite simplificar/actualizar código
        self.inp_t_codigo.set_focus()
        dialog.exec()

    def _save_tipo(self, dialog: CustomDialog):
        data = {
            "tipo_tramite_id": self.current_tipo_id,
            "codigo": self.inp_t_codigo.text().strip().upper(),
            "alias": self.inp_t_alias.text().strip() or None,
            "nombre": self.inp_t_nombre.text().strip(),
            "aplica_a": self.cmb_t_aplica.currentText(),
            "etapa_disparador": self.cmb_t_etapa.currentText(),
            "dias_vigencia": int(self.inp_t_vigencia.text().strip() or 0),
            "es_obligatorio": self.chk_t_obligatorio.isChecked(),
            "activo": self.chk_t_activo.isChecked()
        }
        try:
            if self.api_client.connect_via_api:
                payload = {"usuario_id": self.current_user_id, "sesion_id": self.current_sesion_id, "data": data}
                self.api_client.request("POST", "/api/admin/save/satc_tipos_tramite", data=payload)
            else:
                with self.db_connector.get_session() as session:
                    svc = HousingLifeCycleService(session)
                    svc.save_tipo_tramite(data)
                    session.commit()
            QMessageBox.information(self, "Éxito", "Tipo de trámite guardado correctamente.")
            dialog.accept()
            self.refresh_data()
        except Exception as e:
            QMessageBox.critical(self, "Error", str(e))

    # =========================================================================
    # 2. DIALOGOS & ACCIONES: ESTADOS DE TRÁMITE
    # =========================================================================
    def _create_est_tramite_dialog(self, title: str) -> CustomDialog:
        dialog = CustomDialog(title, self)
        dialog.setMinimumWidth(540)

        card = QFrame(dialog)
        card.setObjectName("card_est_tramite")
        card.setStyleSheet(f"""
            QFrame#card_est_tramite {{
                background-color: {Colors.SLATE_50};
                border: 1px solid {Colors.SLATE_200};
                border-radius: 8px;
                padding: 10px 14px;
            }}
        """)
        lay = QVBoxLayout(card)
        lay.setContentsMargins(0, 0, 0, 0)
        lay.setSpacing(8)

        lbl = CustomLabel("🚦 ESTADO DE TRÁMITE OPERATIVO", variant="subheader")
        lbl.setStyleSheet(f"font-size: 11px; font-weight: bold; color: {Colors.TEXT_LIGHT_PRIMARY}; margin-bottom: 2px;")
        lay.addWidget(lbl)

        form = QFormLayout()
        form.setSpacing(8)
        form.setContentsMargins(0, 0, 0, 0)

        self.inp_et_codigo = CustomInput("Código único (ej. INGRESADO_MUNICIPIO)", parent=card)
        self.inp_et_codigo.setMaxLength(50)
        self.inp_et_codigo.textEdited.connect(lambda t: self.inp_et_codigo.setText(t.upper()))
        self.inp_et_codigo.text = self.inp_et_codigo.text
        self.inp_et_codigo.set_text = self.inp_et_codigo.setText
        self.inp_et_codigo.set_focus = self.inp_et_codigo.setFocus

        self.inp_et_nombre = CustomInput("Nombre descriptivo del estado", parent=card)
        self.inp_et_nombre.setMaxLength(100)
        self.inp_et_nombre.text = self.inp_et_nombre.text
        self.inp_et_nombre.set_text = self.inp_et_nombre.setText

        self.cmb_et_fase = CustomComboBox(parent=card)
        self.cmb_et_fase.add_items(["PREVENTIVA", "DOCUMENTAL", "OPERATIVA", "AUTOMATIZADA", "FINANCIERA", "CIERRE"])

        self.chk_et_bloqueante = CustomCheckBox("Semáforo Bloqueante (Impide continuar sin solventar)", parent=card)
        self.chk_et_bloqueante.setChecked(True)

        self.chk_et_terminal = CustomCheckBox("Estado Terminal (Concluye el ciclo de vida del trámite)", parent=card)
        self.chk_et_terminal.setChecked(False)

        self.chk_et_activo = CustomCheckBox("Activo", parent=card)
        self.chk_et_activo.setChecked(True)

        form.addRow("Código *:", self.inp_et_codigo)
        form.addRow("Nombre *:", self.inp_et_nombre)
        form.addRow("Fase Operativa *:", self.cmb_et_fase)
        form.addRow("", self.chk_et_bloqueante)
        form.addRow("", self.chk_et_terminal)
        form.addRow("", self.chk_et_activo)

        lay.addLayout(form)
        dialog.add_widget(card)

        def _val():
            c = self.inp_et_codigo.text().strip()
            n = self.inp_et_nombre.text().strip()
            dialog.btn_save.setEnabled(bool(c and n))

        self.inp_et_codigo.textChanged.connect(_val)
        self.inp_et_nombre.textChanged.connect(_val)
        _val()

        if not self.can_edit:
            dialog.btn_save.setVisible(False)
            self.inp_et_codigo.setReadOnly(True)
            self.inp_et_nombre.setReadOnly(True)

        dialog.btn_save.clicked.disconnect()
        dialog.btn_save.clicked.connect(lambda: self._save_est_tramite(dialog))
        return dialog

    def _on_new_est_tramite(self):
        self.current_est_tramite_id = None
        dialog = self._create_est_tramite_dialog("Nuevo Estado de Trámite")
        self.inp_et_codigo.set_focus()
        dialog.exec()

    def _on_edit_est_tramite(self, data: dict):
        self.current_est_tramite_id = data.get("estado_tramite_id")
        dialog = self._create_est_tramite_dialog(f"Editar Estado: {data.get('codigo')}")
        self.inp_et_codigo.set_text(data.get("codigo", "") or "")
        self.inp_et_nombre.set_text(data.get("nombre", "") or "")
        self.cmb_et_fase.set_current_text(data.get("fase", "OPERATIVA"))
        self.chk_et_bloqueante.setChecked(bool(data.get("es_bloqueante", True)))
        self.chk_et_terminal.setChecked(bool(data.get("es_terminal", False)))
        self.chk_et_activo.setChecked(bool(data.get("activo", True)))
        self.inp_et_codigo.setReadOnly(True)
        self.inp_et_nombre.set_focus()
        dialog.exec()

    def _save_est_tramite(self, dialog: CustomDialog):
        data = {
            "estado_tramite_id": self.current_est_tramite_id,
            "codigo": self.inp_et_codigo.text().strip().upper(),
            "nombre": self.inp_et_nombre.text().strip(),
            "fase": self.cmb_et_fase.currentText(),
            "es_bloqueante": self.chk_et_bloqueante.isChecked(),
            "es_terminal": self.chk_et_terminal.isChecked(),
            "activo": self.chk_et_activo.isChecked()
        }
        try:
            if self.api_client.connect_via_api:
                payload = {"usuario_id": self.current_user_id, "sesion_id": self.current_sesion_id, "data": data}
                self.api_client.request("POST", "/api/admin/save/satc_estados_tramite", data=payload)
            else:
                with self.db_connector.get_session() as session:
                    svc = HousingLifeCycleService(session)
                    svc.save_estado_tramite(data)
                    session.commit()
            QMessageBox.information(self, "Éxito", "Estado de trámite guardado correctamente.")
            dialog.accept()
            self.refresh_data()
        except Exception as e:
            QMessageBox.critical(self, "Error", str(e))

    # =========================================================================
    # 3. DIALOGOS & ACCIONES: ESTADOS DE VIVIENDA
    # =========================================================================
    def _create_est_vivienda_dialog(self, title: str) -> CustomDialog:
        dialog = CustomDialog(title, self)
        dialog.setMinimumWidth(540)

        card = QFrame(dialog)
        card.setObjectName("card_est_vivienda")
        card.setStyleSheet(f"""
            QFrame#card_est_vivienda {{
                background-color: {Colors.SLATE_50};
                border: 1px solid {Colors.SLATE_200};
                border-radius: 8px;
                padding: 10px 14px;
            }}
        """)
        lay = QVBoxLayout(card)
        lay.setContentsMargins(0, 0, 0, 0)
        lay.setSpacing(8)

        lbl = CustomLabel("🏠 CICLO FÍSICO-COMERCIAL DE LA VIVIENDA", variant="subheader")
        lbl.setStyleSheet(f"font-size: 11px; font-weight: bold; color: {Colors.TEXT_LIGHT_PRIMARY}; margin-bottom: 2px;")
        lay.addWidget(lbl)

        form = QFormLayout()
        form.setSpacing(8)
        form.setContentsMargins(0, 0, 0, 0)

        self.inp_ev_codigo = CustomInput("Código único (ej. EN_INDIVIDUALIZACION)", parent=card)
        self.inp_ev_codigo.setMaxLength(50)
        self.inp_ev_codigo.textEdited.connect(lambda t: self.inp_ev_codigo.setText(t.upper()))
        self.inp_ev_codigo.text = self.inp_ev_codigo.text
        self.inp_ev_codigo.set_text = self.inp_ev_codigo.setText
        self.inp_ev_codigo.set_focus = self.inp_ev_codigo.setFocus

        self.inp_ev_nombre = CustomInput("Nombre del estado", parent=card)
        self.inp_ev_nombre.setMaxLength(100)
        self.inp_ev_nombre.text = self.inp_ev_nombre.text
        self.inp_ev_nombre.set_text = self.inp_ev_nombre.setText

        self.inp_ev_desc = CustomInput("Descripción u observaciones", parent=card)
        self.inp_ev_desc.setMaxLength(250)
        self.inp_ev_desc.text = self.inp_ev_desc.text
        self.inp_ev_desc.set_text = self.inp_ev_desc.setText

        self.inp_ev_orden = CustomInput("Orden secuencial de flujo (1 a 10)", parent=card)
        self.inp_ev_orden.setMaxLength(3)
        self.inp_ev_orden.text = self.inp_ev_orden.text
        self.inp_ev_orden.set_text = self.inp_ev_orden.setText

        self.chk_ev_venta = CustomCheckBox("Habilita Permiso de Venta", parent=card)
        self.chk_ev_escrituracion = CustomCheckBox("Habilita Permiso de Escrituración", parent=card)
        self.chk_ev_activo = CustomCheckBox("Activo", parent=card)
        self.chk_ev_activo.setChecked(True)

        form.addRow("Código *:", self.inp_ev_codigo)
        form.addRow("Nombre *:", self.inp_ev_nombre)
        form.addRow("Descripción:", self.inp_ev_desc)
        form.addRow("Orden Flujo *:", self.inp_ev_orden)
        form.addRow("", self.chk_ev_venta)
        form.addRow("", self.chk_ev_escrituracion)
        form.addRow("", self.chk_ev_activo)

        lay.addLayout(form)
        dialog.add_widget(card)

        def _val():
            c = self.inp_ev_codigo.text().strip()
            n = self.inp_ev_nombre.text().strip()
            dialog.btn_save.setEnabled(bool(c and n))

        self.inp_ev_codigo.textChanged.connect(_val)
        self.inp_ev_nombre.textChanged.connect(_val)
        _val()

        if not self.can_edit:
            dialog.btn_save.setVisible(False)
            self.inp_ev_codigo.setReadOnly(True)
            self.inp_ev_nombre.setReadOnly(True)
            self.inp_ev_orden.setReadOnly(True)

        dialog.btn_save.clicked.disconnect()
        dialog.btn_save.clicked.connect(lambda: self._save_est_vivienda(dialog))
        return dialog

    def _on_new_est_vivienda(self):
        self.current_est_vivienda_id = None
        dialog = self._create_est_vivienda_dialog("Nuevo Estado de Vivienda")
        self.inp_ev_orden.set_text("1")
        self.inp_ev_codigo.set_focus()
        dialog.exec()

    def _on_edit_est_vivienda(self, data: dict):
        self.current_est_vivienda_id = data.get("estado_vivienda_id")
        dialog = self._create_est_vivienda_dialog(f"Editar Estado: {data.get('codigo')}")
        self.inp_ev_codigo.set_text(data.get("codigo", "") or "")
        self.inp_ev_nombre.set_text(data.get("nombre", "") or "")
        self.inp_ev_desc.set_text(data.get("descripcion", "") or "")
        self.inp_ev_orden.set_text(str(data.get("orden_flujo", 1) or 1))
        self.chk_ev_venta.setChecked(bool(data.get("permite_venta", False)))
        self.chk_ev_escrituracion.setChecked(bool(data.get("permite_escrituracion", False)))
        self.chk_ev_activo.setChecked(bool(data.get("activo", True)))
        self.inp_ev_codigo.setReadOnly(True)
        self.inp_ev_nombre.set_focus()
        dialog.exec()

    def _save_est_vivienda(self, dialog: CustomDialog):
        data = {
            "estado_vivienda_id": self.current_est_vivienda_id,
            "codigo": self.inp_ev_codigo.text().strip().upper(),
            "nombre": self.inp_ev_nombre.text().strip(),
            "descripcion": self.inp_ev_desc.text().strip(),
            "orden_flujo": int(self.inp_ev_orden.text().strip() or 1),
            "permite_venta": self.chk_ev_venta.isChecked(),
            "permite_escrituracion": self.chk_ev_escrituracion.isChecked(),
            "activo": self.chk_ev_activo.isChecked()
        }
        try:
            if self.api_client.connect_via_api:
                payload = {"usuario_id": self.current_user_id, "sesion_id": self.current_sesion_id, "data": data}
                self.api_client.request("POST", "/api/admin/save/satc_estados_vivienda", data=payload)
            else:
                with self.db_connector.get_session() as session:
                    svc = HousingLifeCycleService(session)
                    svc.save_estado_vivienda(data)
                    session.commit()
            QMessageBox.information(self, "Éxito", "Estado de vivienda guardado correctamente.")
            dialog.accept()
            self.refresh_data()
        except Exception as e:
            QMessageBox.critical(self, "Error", str(e))

    # =========================================================================
    # 4. DIALOGOS & ACCIONES: REGÍMENES MZ
    # =========================================================================
    def _create_regimen_dialog(self, title: str) -> CustomDialog:
        dialog = CustomDialog(title, self)
        dialog.setMinimumWidth(580)

        card = QFrame(dialog)
        card.setObjectName("card_regimen")
        card.setStyleSheet(f"""
            QFrame#card_regimen {{
                background-color: {Colors.SLATE_50};
                border: 1px solid {Colors.SLATE_200};
                border-radius: 8px;
                padding: 10px 14px;
            }}
        """)
        lay = QVBoxLayout(card)
        lay.setContentsMargins(0, 0, 0, 0)
        lay.setSpacing(8)

        lbl = CustomLabel("🏢 MACRO-LOTE DE RÉGIMEN MZ", variant="subheader")
        lbl.setStyleSheet(f"font-size: 11px; font-weight: bold; color: {Colors.TEXT_LIGHT_PRIMARY}; margin-bottom: 2px;")
        lay.addWidget(lbl)

        form = QFormLayout()
        form.setSpacing(8)
        form.setContentsMargins(0, 0, 0, 0)

        self.cmb_r_desarrollo = CustomComboBox(parent=card)
        self.cmb_r_desarrollo.addItem("-- Selecciona un Desarrollo --", None)
        for des in self.desarrollos_list:
            self.cmb_r_desarrollo.addItem(des.get("nombre", ""), des.get("desarrollo_id"))

        self.inp_r_nombre = CustomInput("Nombre del régimen (ej. REGIMEN AURA NOVA PLAYA)", parent=card)
        self.inp_r_nombre.setMaxLength(200)
        self.inp_r_nombre.text = self.inp_r_nombre.text
        self.inp_r_nombre.set_text = self.inp_r_nombre.setText
        self.inp_r_nombre.set_focus = self.inp_r_nombre.setFocus

        self.inp_r_sm = CustomInput("Super Manzana (ej. SM-115 / Opcional)", parent=card)
        self.inp_r_sm.setMaxLength(50)
        self.inp_r_sm.text = self.inp_r_sm.text
        self.inp_r_sm.set_text = self.inp_r_sm.setText

        self.inp_r_mz = CustomInput("Manzana (ej. MZ-01)", parent=card)
        self.inp_r_mz.setMaxLength(50)
        self.inp_r_mz.text = self.inp_r_mz.text
        self.inp_r_mz.set_text = self.inp_r_mz.setText

        self.inp_r_lt = CustomInput("Lote matriz (ej. LT-01)", parent=card)
        self.inp_r_lt.setMaxLength(50)
        self.inp_r_lt.text = self.inp_r_lt.text
        self.inp_r_lt.set_text = self.inp_r_lt.setText

        self.inp_r_upes = CustomInput("Total de UPEs autorizadas (ej. 577)", parent=card)
        self.inp_r_upes.setMaxLength(6)
        self.inp_r_upes.text = self.inp_r_upes.text
        self.inp_r_upes.set_text = self.inp_r_upes.setText

        self.inp_r_licencia = CustomInput("Licencia de construcción matriz (Opcional)", parent=card)
        self.inp_r_licencia.setMaxLength(100)
        self.inp_r_licencia.text = self.inp_r_licencia.text
        self.inp_r_licencia.set_text = self.inp_r_licencia.setText

        self.inp_r_clave = CustomInput("Clave catastral matriz (Opcional)", parent=card)
        self.inp_r_clave.setMaxLength(100)
        self.inp_r_clave.text = self.inp_r_clave.text
        self.inp_r_clave.set_text = self.inp_r_clave.setText

        self.chk_r_activo = CustomCheckBox("Activo", parent=card)
        self.chk_r_activo.setChecked(True)

        form.addRow("Desarrollo *:", self.cmb_r_desarrollo)
        form.addRow("Nombre Régimen *:", self.inp_r_nombre)
        form.addRow("S.M.:", self.inp_r_sm)
        form.addRow("Manzana *:", self.inp_r_mz)
        form.addRow("Lote Matriz *:", self.inp_r_lt)
        form.addRow("Total UPEs *:", self.inp_r_upes)
        form.addRow("Licencia Matriz:", self.inp_r_licencia)
        form.addRow("Clave Matriz:", self.inp_r_clave)
        form.addRow("", self.chk_r_activo)

        lay.addLayout(form)
        dialog.add_widget(card)

        def _val():
            d = self.cmb_r_desarrollo.currentData()
            n = self.inp_r_nombre.text().strip()
            mz = self.inp_r_mz.text().strip()
            lt = self.inp_r_lt.text().strip()
            u = self.inp_r_upes.text().strip()
            dialog.btn_save.setEnabled(bool(d and n and mz and lt and u))

        self.cmb_r_desarrollo.currentIndexChanged.connect(_val)
        self.inp_r_nombre.textChanged.connect(_val)
        self.inp_r_mz.textChanged.connect(_val)
        self.inp_r_lt.textChanged.connect(_val)
        self.inp_r_upes.textChanged.connect(_val)
        _val()

        if not self.can_edit:
            dialog.btn_save.setVisible(False)
            self.inp_r_nombre.setReadOnly(True)
            self.inp_r_mz.setReadOnly(True)
            self.inp_r_lt.setReadOnly(True)
            self.inp_r_upes.setReadOnly(True)

        dialog.btn_save.clicked.disconnect()
        dialog.btn_save.clicked.connect(lambda: self._save_regimen(dialog))
        return dialog

    def _on_new_regimen(self):
        self.current_regimen_id = None
        dialog = self._create_regimen_dialog("Nuevo Régimen mz")
        self.cmb_r_desarrollo.setCurrentIndex(0)
        self.inp_r_upes.set_text("1")
        self.inp_r_nombre.set_focus()
        dialog.exec()

    def _on_edit_regimen(self, data: dict):
        self.current_regimen_id = data.get("regimen_id")
        dialog = self._create_regimen_dialog(f"Editar Régimen: {data.get('nombre_regimen')}")
        self.inp_r_nombre.set_text(data.get("nombre_regimen", "") or "")
        self.inp_r_sm.set_text(data.get("sm", "") or "")
        self.inp_r_mz.set_text(data.get("mz", "") or "")
        self.inp_r_lt.set_text(data.get("lt", "") or "")
        self.inp_r_upes.set_text(str(data.get("total_upes", 1) or 1))
        self.inp_r_licencia.set_text(data.get("licencia_construccion_matriz", "") or "")
        self.inp_r_clave.set_text(data.get("clave_catastral_matriz", "") or "")
        self.chk_r_activo.setChecked(bool(data.get("activo", True)))

        des_id = data.get("desarrollo_id")
        if des_id:
            idx = self.cmb_r_desarrollo.findData(des_id)
            if idx >= 0:
                self.cmb_r_desarrollo.setCurrentIndex(idx)

        self.inp_r_nombre.set_focus()
        dialog.exec()

    def _save_regimen(self, dialog: CustomDialog):
        des_id = self.cmb_r_desarrollo.currentData()
        if not des_id:
            QMessageBox.warning(self, "Validación", "Debe seleccionar un desarrollo.")
            return

        data = {
            "regimen_id": self.current_regimen_id,
            "desarrollo_id": int(des_id),
            "nombre_regimen": self.inp_r_nombre.text().strip(),
            "sm": self.inp_r_sm.text().strip() or None,
            "mz": self.inp_r_mz.text().strip(),
            "lt": self.inp_r_lt.text().strip(),
            "total_upes": int(self.inp_r_upes.text().strip() or 1),
            "licencia_construccion_matriz": self.inp_r_licencia.text().strip() or None,
            "clave_catastral_matriz": self.inp_r_clave.text().strip() or None,
            "activo": self.chk_r_activo.isChecked()
        }
        try:
            if self.api_client.connect_via_api:
                payload = {"usuario_id": self.current_user_id, "sesion_id": self.current_sesion_id, "data": data}
                self.api_client.request("POST", "/api/admin/save/satc_regimenes", data=payload)
            else:
                with self.db_connector.get_session() as session:
                    svc = HousingLifeCycleService(session)
                    svc.save_regimen_condominio(data)
                    session.commit()
            QMessageBox.information(self, "Éxito", "Régimen guardado correctamente.")
            dialog.accept()
            self.refresh_data()
        except Exception as e:
            QMessageBox.critical(self, "Error", str(e))

    # =========================================================================
    # 5. DIALOGOS & ACCIONES: REGLAS Y FORMATOS MUNICIPALES (Metadata-Driven)
    # =========================================================================
    def _create_formato_dialog(self, title: str) -> CustomDialog:
        dialog = CustomDialog(title, self)
        dialog.setMinimumWidth(620)

        card = QFrame(dialog)
        card.setObjectName("card_formato")
        card.setStyleSheet(f"""
            QFrame#card_formato {{
                background-color: {Colors.SLATE_50};
                border: 1px solid {Colors.SLATE_200};
                border-radius: 8px;
                padding: 10px 14px;
            }}
        """)
        lay = QVBoxLayout(card)
        lay.setContentsMargins(0, 0, 0, 0)
        lay.setSpacing(8)

        lbl = CustomLabel("⚙️ REGLA ESTRUCTURAL DE EXTRACCIÓN Y VALIDACIÓN", variant="subheader")
        lbl.setStyleSheet(f"font-size: 11px; font-weight: bold; color: {Colors.TEXT_LIGHT_PRIMARY}; margin-bottom: 2px;")
        lay.addWidget(lbl)

        form = QFormLayout()
        form.setSpacing(8)

        self.cmb_f_municipio = CustomComboBox(dialog)
        self.cmb_f_municipio.setMinimumHeight(32)
        form.addRow(CustomLabel("Municipio Oficial *:", variant="body"), self.cmb_f_municipio)

        self.cmb_f_campo = CustomComboBox(dialog)
        self.cmb_f_campo.setMinimumHeight(32)
        form.addRow(CustomLabel("Campo Canónico *:", variant="body"), self.cmb_f_campo)

        self.inp_f_variante = CustomInput("Ej. Cancún 18 Caracteres", control_size="md", parent=dialog)
        form.addRow(CustomLabel("Nombre Variante *:", variant="body"), self.inp_f_variante)

        self.inp_f_prioridad = CustomInput("1", control_size="md", parent=dialog)
        form.addRow(CustomLabel("Prioridad Cascada *:", variant="body"), self.inp_f_prioridad)

        self.inp_f_regex_extraccion = CustomInput(r"Ej. \b([1-9](?:\d{17}|\d{16}[A-Z]))\b", control_size="md", parent=dialog)
        form.addRow(CustomLabel("Regex Extracción Directa *:", variant="body"), self.inp_f_regex_extraccion)

        self.inp_f_regex_validacion = CustomInput(r"Ej. ^[1-9](?:\d{17}|\d{16}[A-Z])$", control_size="md", parent=dialog)
        form.addRow(CustomLabel("Regex Validación Estricta *:", variant="body"), self.inp_f_regex_validacion)

        # Contenedor horizontal para longitudes
        lay_long = QHBoxLayout()
        self.inp_f_long_min = CustomInput("Min (ej. 15)", control_size="md", parent=dialog)
        self.inp_f_long_max = CustomInput("Max (ej. 18)", control_size="md", parent=dialog)
        lay_long.addWidget(self.inp_f_long_min)
        lay_long.addWidget(self.inp_f_long_max)
        form.addRow(CustomLabel("Rango Longitud:", variant="body"), lay_long)

        self.inp_f_remover = CustomInput(" -/", control_size="md", parent=dialog)
        form.addRow(CustomLabel("Remover Caracteres:", variant="body"), self.inp_f_remover)

        self.inp_f_error = CustomInput("Mensaje informativo en caso de invalidez", control_size="md", parent=dialog)
        form.addRow(CustomLabel("Descripción Error:", variant="body"), self.inp_f_error)

        self.chk_f_sliding = CustomCheckBox("Habilitar Ventana Deslizante (Sliding Window en bloques fusionados de OCR)", parent=dialog)
        form.addRow("", self.chk_f_sliding)

        self.chk_f_activo = CustomCheckBox("Regla Activa en Producción", parent=dialog)
        self.chk_f_activo.setChecked(True)
        form.addRow("", self.chk_f_activo)

        lay.addLayout(form)
        dialog.add_widget(card)

        # Llenar dropdowns
        self.cmb_f_municipio.clear()
        for m in getattr(self, "_muns_list", []):
            self.cmb_f_municipio.addItem(m.get("nombre", ""), m.get("municipio_id"))

        self.cmb_f_campo.clear()
        for c in getattr(self, "_campos_list", []):
            self.cmb_f_campo.addItem(f"{c.get('nombre')} ({c.get('codigo')})", c.get("campo_id"))

        dialog.set_primary_button_text("Guardar Regla")
        dialog.primary_button_clicked.connect(lambda: self._save_formato(dialog))
        return dialog

    def _on_new_formato(self):
        self.current_formato_id = None
        dialog = self._create_formato_dialog("Nueva Regla de Extracción Municipal")
        self.inp_f_variante.set_text("")
        self.inp_f_prioridad.set_text("1")
        self.inp_f_regex_extraccion.set_text("")
        self.inp_f_regex_validacion.set_text("")
        self.inp_f_long_min.set_text("")
        self.inp_f_long_max.set_text("")
        self.inp_f_remover.set_text(" -/")
        self.inp_f_error.set_text("")
        self.chk_f_sliding.setChecked(False)
        self.chk_f_activo.setChecked(True)
        self.inp_f_variante.set_focus()
        dialog.exec()

    def _on_edit_formato(self, data: Dict[str, Any]):
        if not data:
            return
        self.current_formato_id = data.get("formato_id")
        dialog = self._create_formato_dialog("Editar Regla de Extracción Municipal")

        mun_id = data.get("municipio_id")
        if mun_id:
            idx_m = self.cmb_f_municipio.findData(mun_id)
            if idx_m >= 0:
                self.cmb_f_municipio.setCurrentIndex(idx_m)

        cmp_id = data.get("campo_id")
        if cmp_id:
            idx_c = self.cmb_f_campo.findData(cmp_id)
            if idx_c >= 0:
                self.cmb_f_campo.setCurrentIndex(idx_c)

        self.inp_f_variante.set_text(data.get("nombre_variante", "") or "")
        self.inp_f_prioridad.set_text(str(data.get("prioridad", 1) or 1))
        self.inp_f_regex_extraccion.set_text(data.get("regex_extraccion", "") or "")
        self.inp_f_regex_validacion.set_text(data.get("regex_validacion", "") or "")
        self.inp_f_long_min.set_text(str(data.get("longitud_min")) if data.get("longitud_min") is not None else "")
        self.inp_f_long_max.set_text(str(data.get("longitud_max")) if data.get("longitud_max") is not None else "")
        self.inp_f_remover.set_text(data.get("remover_caracteres", " -/") or " -/")
        self.inp_f_error.set_text(data.get("descripcion_error", "") or "")
        self.chk_f_sliding.setChecked(bool(data.get("permite_sliding_window", False)))
        self.chk_f_activo.setChecked(bool(data.get("activo", True)))

        self.inp_f_variante.set_focus()
        dialog.exec()

    def _save_formato(self, dialog: CustomDialog):
        mun_id = self.cmb_f_municipio.currentData()
        cmp_id = self.cmb_f_campo.currentData()
        variante = self.inp_f_variante.text().strip()
        regex_val = self.inp_f_regex_validacion.text().strip()
        regex_ext = self.inp_f_regex_extraccion.text().strip()

        if not mun_id or not cmp_id:
            QMessageBox.warning(self, "Validación", "Debe seleccionar un municipio y un campo canónico.")
            return
        if not variante or not regex_val or not regex_ext:
            QMessageBox.warning(self, "Validación", "Nombre de variante, regex de extracción y regex de validación son obligatorios.")
            return

        try:
            prio = int(self.inp_f_prioridad.text().strip() or 1)
        except ValueError:
            prio = 1

        data = {
            "formato_id": self.current_formato_id,
            "municipio_id": int(mun_id),
            "campo_id": int(cmp_id),
            "nombre_variante": variante,
            "prioridad": prio,
            "regex_extraccion": regex_ext,
            "regex_validacion": regex_val,
            "longitud_min": self.inp_f_long_min.text().strip() or None,
            "longitud_max": self.inp_f_long_max.text().strip() or None,
            "remover_caracteres": self.inp_f_remover.text().strip() or " -/",
            "permite_sliding_window": self.chk_f_sliding.isChecked(),
            "descripcion_error": self.inp_f_error.text().strip() or None,
            "activo": self.chk_f_activo.isChecked()
        }

        try:
            if self.api_client.connect_via_api:
                payload = {"usuario_id": self.current_user_id, "sesion_id": self.current_sesion_id, "data": data}
                self.api_client.request("POST", "/api/admin/save/satc_formatos_municipio", data=payload)
            else:
                from satc.src.storage.satc_repos import MetadataRulesRepository
                from satc.src.services.metadata_extraction_service import MetadataExtractionService
                with self.db_connector.get_session() as session:
                    repo = MetadataRulesRepository(session)
                    repo.save_formato(data)
                    session.commit()
                    MetadataExtractionService.clear_cache()
            QMessageBox.information(self, "Éxito", "Regla de extracción municipal guardada correctamente.")
            dialog.accept()
            self.refresh_data()
        except Exception as e:
            QMessageBox.critical(self, "Error", str(e))

    # =========================================================================
    # CARGA Y REFRESH DE DATOS
    # =========================================================================
    def refresh_data(self):
        try:
            if self.api_client.connect_via_api:
                self.desarrollos_list = self.api_client.request("GET", "/api/admin/data/desarrollos")
                tipos = self.api_client.request("GET", "/api/admin/data/satc_tipos_tramite")
                est_tramite = self.api_client.request("GET", "/api/admin/data/satc_estados_tramite")
                est_vivienda = self.api_client.request("GET", "/api/admin/data/satc_estados_vivienda")
                regimenes = self.api_client.request("GET", "/api/admin/data/satc_regimenes")
                formatos = self.api_client.request("GET", "/api/admin/data/satc_formatos_municipio")
                self._muns_list = self.api_client.request("GET", "/api/admin/data/municipios")
                self._campos_list = self.api_client.request("GET", "/api/admin/data/satc_campos")
            else:
                with self.db_connector.get_session() as session:
                    cat_repo = CatalogoRepository(session)
                    des_objs = cat_repo.get_all_desarrollos()
                    self.desarrollos_list = [
                        {"desarrollo_id": d.desarrollo_id, "nombre": d.nombre, "alias": d.alias}
                        for d in des_objs
                    ]

                    # Municipios para diálogos
                    muns_objs = cat_repo.get_all_municipios()
                    self._muns_list = [
                        {"municipio_id": m.municipio_id, "nombre": m.nombre, "codigo_portal": m.codigo_portal}
                        for m in muns_objs
                    ]

                    satc_svc = HousingLifeCycleService(session)
                    tipos_objs = satc_svc.get_all_tipos_tramite()
                    tipos = [
                        {
                            "tipo_tramite_id": t.tipo_tramite_id,
                            "codigo": t.codigo,
                            "nombre": t.nombre,
                            "aplica_a": t.aplica_a,
                            "etapa_disparador": t.etapa_disparador,
                            "dias_vigencia": t.dias_vigencia,
                            "es_obligatorio": t.es_obligatorio,
                            "activo": t.activo
                        }
                        for t in tipos_objs
                    ]

                    est_tr_objs = satc_svc.get_all_estados_tramite()
                    est_tramite = [
                        {
                            "estado_tramite_id": e.estado_tramite_id,
                            "codigo": e.codigo,
                            "nombre": e.nombre,
                            "fase": e.fase,
                            "es_bloqueante": e.es_bloqueante,
                            "es_terminal": e.es_terminal,
                            "activo": e.activo
                        }
                        for e in est_tr_objs
                    ]

                    est_viv_objs = satc_svc.get_all_estados_vivienda()
                    est_vivienda = [
                        {
                            "estado_vivienda_id": e.estado_vivienda_id,
                            "codigo": e.codigo,
                            "nombre": e.nombre,
                            "descripcion": e.descripcion,
                            "orden_flujo": e.orden_flujo,
                            "permite_venta": e.permite_venta,
                            "permite_escrituracion": e.permite_escrituracion,
                            "activo": e.activo
                        }
                        for e in est_viv_objs
                    ]

                    reg_objs = satc_svc.get_all_regimenes()
                    regimenes = [
                        {
                            "regimen_id": r.regimen_id,
                            "desarrollo_id": r.desarrollo_id,
                            "desarrollo_nombre": r.desarrollo.nombre if r.desarrollo else "N/A",
                            "nombre_regimen": r.nombre_regimen,
                            "sm": r.sm or "",
                            "mz": r.mz,
                            "lt": r.lt,
                            "total_upes": r.total_upes,
                            "upes_generadas": r.upes_generadas,
                            "licencia_construccion_matriz": r.licencia_construccion_matriz or "",
                            "clave_catastral_matriz": r.clave_catastral_matriz or "",
                            "activo": r.activo
                        }
                        for r in reg_objs
                    ]

                    # Metadatos de Extracción SATC
                    from satc.src.storage.satc_repos import MetadataRulesRepository
                    meta_repo = MetadataRulesRepository(session)
                    campos_objs = meta_repo.list_campos()
                    self._campos_list = [
                        {
                            "campo_id": c.campo_id,
                            "codigo": c.codigo,
                            "nombre": c.nombre,
                            "descripcion": c.descripcion or "",
                            "activo": c.activo
                        }
                        for c in campos_objs
                    ]

                    formatos_objs = meta_repo.list_all_formatos(solo_activos=False)
                    formatos = [
                        {
                            "formato_id": f.formato_id,
                            "campo_id": f.campo_id,
                            "campo_nombre": f.campo.nombre if f.campo else "",
                            "campo_codigo": f.campo.codigo if f.campo else "",
                            "municipio_id": f.municipio_id,
                            "municipio_nombre": f.municipio.nombre if f.municipio else "TODOS",
                            "nombre_variante": f.nombre_variante,
                            "regex_extraccion": f.regex_extraccion,
                            "regex_validacion": f.regex_validacion,
                            "longitud_min": f.longitud_min,
                            "longitud_max": f.longitud_max,
                            "remover_caracteres": f.remover_caracteres or "",
                            "permite_sliding_window": f.permite_sliding_window,
                            "descripcion_error": f.descripcion_error or "",
                            "prioridad": f.prioridad,
                            "activo": f.activo
                        }
                        for f in formatos_objs
                    ]

            self.tbl_tipos.populate(tipos)
            self.tbl_est_tramite.populate(est_tramite)
            self.tbl_est_vivienda.populate(est_vivienda)
            self.tbl_regimenes.populate(regimenes)
            if hasattr(self, "tbl_formatos"):
                self.tbl_formatos.populate(formatos)

            # Actualizar dropdown de desarrollos en pestaña de viviendas
            current_v_des = self.cmb_v_desarrollo.currentData()
            self.cmb_v_desarrollo.blockSignals(True)
            self.cmb_v_desarrollo.clear()
            self.cmb_v_desarrollo.addItem("-- Selecciona un Desarrollo --", None)
            for d in self.desarrollos_list:
                self.cmb_v_desarrollo.addItem(d.get("nombre", ""), d.get("desarrollo_id"))
            if current_v_des:
                idx = self.cmb_v_desarrollo.findData(current_v_des)
                if idx >= 0:
                    self.cmb_v_desarrollo.setCurrentIndex(idx)
                else:
                    self.cmb_v_desarrollo.setCurrentIndex(0)
            else:
                self.cmb_v_desarrollo.setCurrentIndex(0)
            self.cmb_v_desarrollo.blockSignals(False)
            self._on_v_desarrollo_changed()

        except Exception as e:
            QMessageBox.critical(self, "Error al cargar catálogos SATC", str(e))

    # =========================================================================
    # 5. HANDLERS: IMPORTACIÓN Y PADRÓN DE VIVIENDAS / UPES
    # =========================================================================
    def _on_v_desarrollo_changed(self):
        """Recarga los regímenes disponibles para el desarrollo preservando la selección si aún pertenece a él."""
        des_id = self.cmb_v_desarrollo.currentData()
        current_reg_id = self.cmb_v_regimen.currentData()

        self.cmb_v_regimen.blockSignals(True)
        self.cmb_v_regimen.clear()
        self.cmb_v_regimen.addItem("-- Seleccione un Régimen --", None)

        if des_id:
            try:
                if self.api_client.connect_via_api:
                    all_regs = self.api_client.request("GET", "/api/admin/data/satc_regimenes")
                    regs = [r for r in all_regs if r.get("desarrollo_id") == des_id and r.get("activo", True)]
                else:
                    with self.db_connector.get_session() as session:
                        svc = HousingLifeCycleService(session)
                        all_regs = svc.get_all_regimenes()
                        regs = [
                            {"regimen_id": r.regimen_id, "nombre_regimen": r.nombre_regimen, "total_upes": r.total_upes}
                            for r in all_regs if r.desarrollo_id == des_id and r.activo
                        ]

                for r in regs:
                    total_u = r.get("total_upes", 0)
                    self.cmb_v_regimen.addItem(f"{r.get('nombre_regimen')} (Capacidad: {total_u} UPEs)", r.get("regimen_id"))

                # Restaurar el régimen previo si sigue existiendo en el nuevo listado
                if current_reg_id:
                    idx_reg = self.cmb_v_regimen.findData(current_reg_id)
                    if idx_reg >= 0:
                        self.cmb_v_regimen.setCurrentIndex(idx_reg)
                    else:
                        self.cmb_v_regimen.setCurrentIndex(0)
                else:
                    self.cmb_v_regimen.setCurrentIndex(0)
            except Exception:
                pass

        self.cmb_v_regimen.blockSignals(False)
        self._on_v_regimen_changed()

    def _on_v_regimen_changed(self):
        """Carga las viviendas de forma asíncrona en segundo plano sin congelar la interfaz."""
        des_id = self.cmb_v_desarrollo.currentData()
        reg_id = self.cmb_v_regimen.currentData()

        # Detener limpiamente cualquier worker de carga en ejecución previa
        if self._v_load_worker and self._v_load_worker.isRunning():
            try:
                self._v_load_worker.finished.disconnect()
            except Exception:
                pass
            self._v_load_worker.quit()
            self._v_load_worker.wait(50)
            self._v_load_worker = None

        if self._v_loading_dialog:
            try:
                self._v_loading_dialog.close()
            except Exception:
                pass
            self._v_loading_dialog = None

        # Reset y limpieza inmediata de estado y tabla
        self._v_parsed_records = []
        self.btn_v_ejecutar.setEnabled(False)
        self.btn_v_exportar_datos.setEnabled(False)
        self.lbl_vkpi_archivo.setText("Archivo: Ninguno")
        self.lbl_vkpi_total_arch.setText("En Archivo: 0")
        self.lbl_vkpi_nuevos.setText("Nuevos / Vinculables: 0")
        self.lbl_vkpi_duplicados.setText("Duplicados / Sin Coincidencia: 0")
        if hasattr(self, "kpi_alertas_frame"):
            self.lbl_vkpi_alertas.setText("Con Advertencia: 0")
            self.kpi_alertas_frame.hide()
        if hasattr(self, "kpi_faltantes_frame"):
            self.lbl_vkpi_faltantes.setText("Faltantes: 0")
            self.kpi_faltantes_frame.hide()
        self.search_viviendas_input.blockSignals(True)
        self.search_viviendas_input.clear()
        self.search_viviendas_input.blockSignals(False)

        # Limpiar la tabla de forma atómica para que el usuario perciba el cambio de inmediato
        self.tbl_viviendas.setRowCount(0)

        # Si no se ha seleccionado desarrollo o régimen, mantener la tabla limpia en reposo
        if not des_id or reg_id is None:
            self.lbl_vkpi_bd.setText("EnBD: 0")
            self.lbl_vkpi_regimen_info.setText("Total UPE: N/A")
            self.lbl_vkpi_con_rpp.setText("Con Folio: 0")
            self.lbl_vkpi_con_catastro.setText("Con CC: 0")
            self.action_state_changed.emit()
            return

        # Desplegar estado de carga visual en el encabezado
        self.lbl_vkpi_archivo.setText("Consultando viviendas del régimen...")
        self.lbl_vkpi_archivo.setStyleSheet(f"font-size: 12px; color: {Colors.ACCENT}; font-weight: bold;")
        self.lbl_vkpi_bd.setText("Cargando...")

        # Iniciar worker en background
        self._v_load_worker = ViviendasLoadWorker(
            api_client=self.api_client,
            db_connector=self.db_connector,
            desarrollo_id=des_id,
            regimen_id=reg_id
        )
        self._v_load_worker.finished.connect(self._on_v_load_finished)
        self._v_load_worker.start()

    def _on_v_load_finished(self, success: bool, result_data: dict, error_msg: str):
        """Puebla la tabla de viviendas y las métricas una vez culminada la consulta asíncrona."""
        self.lbl_vkpi_archivo.setText("Archivo: Ninguno")
        self.lbl_vkpi_archivo.setStyleSheet(f"font-size: 12px; color: {Colors.TEXT_LIGHT_MUTED}; font-weight: normal;")
        if hasattr(self, "kpi_faltantes_frame"):
            self.kpi_faltantes_frame.hide()
            self.lbl_vkpi_faltantes.setText("Faltantes: 0")
        if hasattr(self, "cmb_v_filtro_estado"):
            self.cmb_v_filtro_estado.blockSignals(True)
            self.cmb_v_filtro_estado.clear()
            self.cmb_v_filtro_estado.addItem("Mostrar: Todos", "TODOS")
            self.cmb_v_filtro_estado.blockSignals(False)
        self._highlight_kpi_active("TODOS")

        if not success:
            QMessageBox.critical(self, "Error al Cargar Viviendas", f"No se pudieron cargar las viviendas: {error_msg}")
            return

        viviendas = result_data.get("viviendas", [])
        tot_u = result_data.get("tot_u", 0)
        reg_id = result_data.get("regimen_id")

        # Comprobar que el resultado corresponda al régimen actualmente seleccionado
        current_reg_id = self.cmb_v_regimen.currentData()
        if current_reg_id != reg_id:
            return

        self._v_bd_count = len(viviendas)
        self._v_tot_upes = tot_u

        if reg_id:
            self.lbl_vkpi_regimen_info.setText(f"Total UPE: {tot_u}")
        else:
            self.lbl_vkpi_regimen_info.setText("Sin Régimen (Abierto)")

        self.lbl_vkpi_bd.setText(f"EnBD: {len(viviendas)}")

        # Verificación preventiva si el usuario tenía seleccionado 'INICIAL'
        if self.cmb_v_modo.currentData() == "INICIAL" and tot_u > 0 and len(viviendas) >= tot_u:
            self._on_v_modo_changed()

        con_rpp = sum(1 for v in viviendas if v.get("folio_electronico"))
        con_cat = sum(1 for v in viviendas if v.get("clave_catastral"))
        self.lbl_vkpi_con_rpp.setText(f"Con Folio: {con_rpp}")
        self.lbl_vkpi_con_catastro.setText(f"Con CC: {con_cat}")

        self.tbl_viviendas.setUpdatesEnabled(False)
        try:
            self.tbl_viviendas.setRowCount(len(viviendas))
            for idx, v in enumerate(viviendas):
                item_st = QTableWidgetItem("REGISTRADO")
                item_st.setForeground(Qt.darkGreen)
                self.tbl_viviendas.setItem(idx, 0, item_st)
                self.tbl_viviendas.setItem(idx, 1, QTableWidgetItem(str(v.get("regimen_nombre") or "")))
                self.tbl_viviendas.setItem(idx, 2, QTableWidgetItem(str(v.get("numero_oficial") or "")))
                self.tbl_viviendas.setItem(idx, 3, QTableWidgetItem(str(v.get("interior") or "")))
                self.tbl_viviendas.setItem(idx, 4, QTableWidgetItem(str(v.get("edificio") or "")))
                self.tbl_viviendas.setItem(idx, 5, QTableWidgetItem(str(v.get("nivel") or "")))
                self.tbl_viviendas.setItem(idx, 6, QTableWidgetItem(str(v.get("destino") or "HABITACIONAL")))
                self.tbl_viviendas.setItem(idx, 7, QTableWidgetItem(str(v.get("espacio") or "VIVIENDA")))
                self.tbl_viviendas.setItem(idx, 8, QTableWidgetItem(str(v.get("folio_electronico") or "")))
                self.tbl_viviendas.setItem(idx, 9, QTableWidgetItem(str(v.get("clave_catastral") or "")))
                self.tbl_viviendas.setItem(idx, 10, QTableWidgetItem(str(v.get("padron") or "")))
                self.tbl_viviendas.setItem(idx, 11, QTableWidgetItem(str(v.get("calle") or "")))
                self.tbl_viviendas.setItem(idx, 12, QTableWidgetItem(str(v.get("descripcion") or "")))
        finally:
            self.tbl_viviendas.setUpdatesEnabled(True)

        self._v_db_viviendas = viviendas
        self.btn_v_exportar_datos.setEnabled(len(viviendas) > 0)
        self.action_state_changed.emit()

    def _render_viviendas_table_from_cache(self, viviendas: list):
        """Puebla la tabla de viviendas a partir de la caché en memoria evitando consultas asíncronas y parpadeos."""
        self.tbl_viviendas.setUpdatesEnabled(False)
        try:
            self.tbl_viviendas.setRowCount(len(viviendas))
            for idx, v in enumerate(viviendas):
                item_st = QTableWidgetItem("REGISTRADO")
                item_st.setForeground(Qt.darkGreen)
                self.tbl_viviendas.setItem(idx, 0, item_st)
                self.tbl_viviendas.setItem(idx, 1, QTableWidgetItem(str(v.get("regimen_nombre") or "")))
                self.tbl_viviendas.setItem(idx, 2, QTableWidgetItem(str(v.get("numero_oficial") or "")))
                self.tbl_viviendas.setItem(idx, 3, QTableWidgetItem(str(v.get("interior") or "")))
                self.tbl_viviendas.setItem(idx, 4, QTableWidgetItem(str(v.get("edificio") or "")))
                self.tbl_viviendas.setItem(idx, 5, QTableWidgetItem(str(v.get("nivel") or "")))
                self.tbl_viviendas.setItem(idx, 6, QTableWidgetItem(str(v.get("destino") or "HABITACIONAL")))
                self.tbl_viviendas.setItem(idx, 7, QTableWidgetItem(str(v.get("espacio") or "VIVIENDA")))
                self.tbl_viviendas.setItem(idx, 8, QTableWidgetItem(str(v.get("folio_electronico") or "")))
                self.tbl_viviendas.setItem(idx, 9, QTableWidgetItem(str(v.get("clave_catastral") or "")))
                self.tbl_viviendas.setItem(idx, 10, QTableWidgetItem(str(v.get("padron") or "")))
                self.tbl_viviendas.setItem(idx, 11, QTableWidgetItem(str(v.get("calle") or "")))
                self.tbl_viviendas.setItem(idx, 12, QTableWidgetItem(str(v.get("descripcion") or "")))
        finally:
            self.tbl_viviendas.setUpdatesEnabled(True)

    def _on_kpi_clicked(self, kpi_key: str):
        """Maneja el clic interactivo en una pastilla KPI para filtrar la tabla con efecto toggle."""
        modo = self.cmb_v_modo.currentData() if hasattr(self, "cmb_v_modo") else "INICIAL"
        
        # Mapeo de KPI a código de estado en la tabla
        target_status = "TODOS"
        if kpi_key == "VINCULABLES":
            target_status = "VINCULABLE" if modo in ("RPP", "CATASTRO") else "NUEVO"
        elif kpi_key == "RECHAZADOS":
            target_status = "NO ENCONTRADO" if modo in ("RPP", "CATASTRO") else "DUPLICADO"
        elif kpi_key == "ALERTAS":
            target_status = "ALERTAS"
        elif kpi_key == "FALTANTES":
            target_status = "FALTANTE EN ARCHIVO"
        elif kpi_key == "TOTAL":
            target_status = "TODOS"

        current_selected = self.cmb_v_filtro_estado.currentData() if hasattr(self, "cmb_v_filtro_estado") else "TODOS"

        # Efecto Toggle: Si ya está seleccionado ese filtro, regresar a TODOS
        if target_status != "TODOS" and current_selected == target_status:
            target_status = "TODOS"

        # Sincronizar el selector de estado
        if hasattr(self, "cmb_v_filtro_estado"):
            found_idx = -1
            for i in range(self.cmb_v_filtro_estado.count()):
                if self.cmb_v_filtro_estado.itemData(i) == target_status:
                    found_idx = i
                    break
            if found_idx != -1:
                self.cmb_v_filtro_estado.setCurrentIndex(found_idx)
            else:
                self._apply_viviendas_filters(custom_status=target_status)
        else:
            self._apply_viviendas_filters(custom_status=target_status)

        self._highlight_kpi_active(target_status)

    def _highlight_kpi_active(self, active_status: str):
        """Actualiza el resalte visual de la pastilla KPI activa acorde a tokens de color."""
        if not hasattr(self, "kpi_containers"):
            return

        modo = self.cmb_v_modo.currentData() if hasattr(self, "cmb_v_modo") else "INICIAL"
        vinculables_code = "VINCULABLE" if modo in ("RPP", "CATASTRO") else "NUEVO"
        rechazados_code = "NO ENCONTRADO" if modo in ("RPP", "CATASTRO") else "DUPLICADO"

        key_mapping = {
            "TOTAL": active_status == "TODOS",
            "VINCULABLES": active_status == vinculables_code,
            "RECHAZADOS": active_status == rechazados_code,
            "ALERTAS": active_status == "ALERTAS",
            "FALTANTES": active_status == "FALTANTE EN ARCHIVO"
        }

        for k, is_active in key_mapping.items():
            info = self.kpi_containers.get(k)
            if not info:
                continue
            container = info["container"]
            dot_color = info["dot_color"]

            if is_active and active_status != "TODOS":
                container.setStyleSheet(f"""
                    QFrame#{container.objectName()} {{
                        background-color: {Colors.NEUTRAL_BG};
                        border: 1px solid {dot_color};
                        border-radius: 4px;
                        padding: 2px 4px;
                    }}
                """)
            else:
                container.setStyleSheet(f"""
                    QFrame#{container.objectName()} {{
                        background-color: transparent;
                        border: 1px solid transparent;
                        border-radius: 4px;
                        padding: 2px 4px;
                    }}
                    QFrame#{container.objectName()}:hover {{
                        background-color: {Colors.NEUTRAL_BG};
                        border: 1px solid {Colors.BORDER_LIGHT};
                    }}
                """)

    def _on_filtro_estado_changed(self, index: int):
        """Filtra las filas de la tabla según el estado seleccionado y resalta la pastilla correspondiente."""
        self._apply_viviendas_filters()
        selected_status = self.cmb_v_filtro_estado.currentData() if hasattr(self, "cmb_v_filtro_estado") else "TODOS"
        self._highlight_kpi_active(selected_status)

    def _on_search_viviendas_text_changed(self, text: str):
        """Filtra reactivamente las filas visibles en la tabla de viviendas en base al texto buscado y al estado."""
        self._apply_viviendas_filters()

    def _apply_viviendas_filters(self, custom_status: Optional[str] = None):
        """Aplica simultáneamente el filtro de estado y el filtro de texto de búsqueda."""
        search_term = self.search_viviendas_input.text().strip().lower()
        selected_status = custom_status if custom_status is not None else (
            self.cmb_v_filtro_estado.currentData() if hasattr(self, "cmb_v_filtro_estado") else "TODOS"
        )

        self.tbl_viviendas.setUpdatesEnabled(False)
        try:
            for row in range(self.tbl_viviendas.rowCount()):
                item_st = self.tbl_viviendas.item(row, 0)
                row_status = item_st.text().strip() if item_st else ""
                item_desc = self.tbl_viviendas.item(row, 12)  # Columna Notas
                desc_text = item_desc.text().strip() if item_desc else ""

                # 1. Filtro por Estado / Alerta
                if selected_status and selected_status != "TODOS":
                    if selected_status == "ALERTAS":
                        # Mostrar exclusivamente filas que tengan advertencia de inconsistencia
                        if "⚠️" not in desc_text and "ERROR EN TEXTO" not in desc_text:
                            self.tbl_viviendas.setRowHidden(row, True)
                            continue
                    else:
                        if row_status != selected_status:
                            self.tbl_viviendas.setRowHidden(row, True)
                            continue

                # 2. Filtro por Texto Libre
                if not search_term:
                    self.tbl_viviendas.setRowHidden(row, False)
                    continue

                row_matches = False
                for col in range(self.tbl_viviendas.columnCount()):
                    item = self.tbl_viviendas.item(row, col)
                    if item and search_term in item.text().lower():
                        row_matches = True
                        break
                self.tbl_viviendas.setRowHidden(row, not row_matches)
        finally:
            self.tbl_viviendas.setUpdatesEnabled(True)

    def _on_v_limpiar(self):
        """Limpia completamente la pantalla, filtros, registros cargados, tabla e indicadores de métricas."""
        self.cmb_v_desarrollo.blockSignals(True)
        self.cmb_v_regimen.blockSignals(True)
        self.cmb_v_modo.blockSignals(True)
        try:
            self.cmb_v_desarrollo.setCurrentIndex(0)
            self.cmb_v_regimen.clear()
            self.cmb_v_regimen.addItem("-- Seleccione un Régimen --", None)
            self.cmb_v_regimen.setCurrentIndex(0)
            self.cmb_v_modo.setCurrentIndex(0)
        finally:
            self.cmb_v_desarrollo.blockSignals(False)
            self.cmb_v_regimen.blockSignals(False)
            self.cmb_v_modo.blockSignals(False)

        self._v_parsed_records = []
        self._v_db_viviendas = []
        self.btn_v_ejecutar.setEnabled(False)
        self.btn_v_exportar_datos.setEnabled(False)
        self.search_viviendas_input.blockSignals(True)
        self.search_viviendas_input.clear()
        self.search_viviendas_input.blockSignals(False)

        if hasattr(self, "cmb_v_filtro_estado"):
            self.cmb_v_filtro_estado.blockSignals(True)
            self.cmb_v_filtro_estado.clear()
            self.cmb_v_filtro_estado.addItem("Mostrar: Todos", "TODOS")
            self.cmb_v_filtro_estado.blockSignals(False)

        # Reset completo de indicadores KPI (Nomenclatura concisa)
        self.lbl_vkpi_archivo.setText("Archivo: Ninguno")
        self.lbl_vkpi_total_arch.setText("En Archivo: 0")
        self.lbl_vkpi_nuevos.setText("Nuevos / Vinculables: 0")
        self.lbl_vkpi_duplicados.setText("Duplicados / Sin Coincidencia: 0")
        if hasattr(self, "kpi_alertas_frame"):
            self.lbl_vkpi_alertas.setText("Con Advertencia: 0")
            self.kpi_alertas_frame.hide()
        if hasattr(self, "lbl_vkpi_faltantes"):
            self.lbl_vkpi_faltantes.setText("Faltantes: 0")
            self.kpi_faltantes_frame.hide()
        self.lbl_vkpi_regimen_info.setText("Total UPE: N/A")
        self.lbl_vkpi_bd.setText("EnBD: 0")
        self.lbl_vkpi_con_rpp.setText("Con Folio: 0")
        self.lbl_vkpi_con_catastro.setText("Con CC: 0")
        self._highlight_kpi_active("TODOS")

        # Limpiar tabla
        self.tbl_viviendas.setRowCount(0)
        self.action_state_changed.emit()

    def _on_v_exportar_datos(self):
        """Descarga a Excel (.xlsx) los registros del archivo analizado o el padrón de viviendas del régimen registrado en BD."""
        records_to_export = []
        origen_tipo = ""

        # Prioridad 1: Si hay registros parsed de un archivo recién cargado
        if getattr(self, "_v_parsed_records", None):
            records_to_export = self._v_parsed_records
            origen_tipo = "Archivo"
        # Prioridad 2: Si no hay archivo cargado pero hay viviendas en BD del régimen seleccionado
        elif getattr(self, "_v_db_viviendas", None):
            records_to_export = self._v_db_viviendas
            origen_tipo = "BD_Regimen"

        if not records_to_export:
            QMessageBox.warning(self, "Sin Datos", "No hay datos de viviendas cargados ni registrados para exportar.")
            return

        modo = self.cmb_v_modo.currentData() or "PADRON"
        nombre_des = "".join(c for c in self.cmb_v_desarrollo.currentText() if c.isalnum() or c in (" ", "_", "-")).strip()
        reg_id = self.cmb_v_regimen.currentData()
        nombre_reg = "".join(c for c in self.cmb_v_regimen.currentText() if c.isalnum() or c in (" ", "_", "-")).strip() if reg_id else ""

        if origen_tipo == "Archivo":
            nombre_sugerido = f"Datos_Extraidos_{modo}_{nombre_des}.xlsx"
            dialog_title = "Descargar Datos Extraídos del Archivo"
            sheet_title = f"Datos_{modo}"
        else:
            nombre_sugerido = f"Padron_Viviendas_{nombre_des}_{nombre_reg}.xlsx"
            dialog_title = "Exportar Padrón de Viviendas del Régimen"
            sheet_title = "Padron_Viviendas"

        file_path, _ = QFileDialog.getSaveFileName(
            self, dialog_title, nombre_sugerido, "Excel Files (*.xlsx)"
        )
        if not file_path:
            return

        try:
            wb = openpyxl.Workbook()
            ws = wb.active
            ws.title = sheet_title[:31]  # Límite máximo de caracteres en nombres de pestañas Excel

            headers = [
                "NUMERO_OFICIAL", "INTERIOR", "EDIFICIO", "NIVEL", "DESTINO", 
                "ESPACIO", "FOLIO_ELECTRONICO", "CLAVE_CATASTRAL", "PADRON", "CALLE", "DESCRIPCION"
            ]
            ws.append(headers)

            # Estilos de encabezado
            header_fill = PatternFill(start_color="1E293B", end_color="1E293B", fill_type="solid")
            header_font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
            border_thin = Border(
                left=Side(style='thin', color='CBD5E1'),
                right=Side(style='thin', color='CBD5E1'),
                top=Side(style='thin', color='CBD5E1'),
                bottom=Side(style='thin', color='CBD5E1')
            )

            for col_num in range(1, len(headers) + 1):
                cell = ws.cell(row=1, column=col_num)
                cell.fill = header_fill
                cell.font = header_font
                cell.alignment = Alignment(horizontal="center", vertical="center")
                cell.border = border_thin

            for r in records_to_export:
                ws.append([
                    r.get("numero_oficial", ""),
                    r.get("interior", ""),
                    r.get("edificio", ""),
                    r.get("nivel", ""),
                    r.get("destino", "HABITACIONAL"),
                    r.get("espacio", "VIVIENDA"),
                    r.get("folio_electronico", ""),
                    r.get("clave_catastral", ""),
                    r.get("padron", ""),
                    r.get("calle", ""),
                    r.get("descripcion", "")
                ])

            # Autoancho de columnas
            for col in ws.columns:
                max_len = max(len(str(cell.value or '')) for cell in col)
                col_letter = get_column_letter(col[0].column)
                ws.column_dimensions[col_letter].width = max(max_len + 4, 12)

            wb.save(file_path)
            QMessageBox.information(
                self, "Exportación Exitosa",
                f"Se han exportado exitosamente {len(records_to_export)} registros a:\n{file_path}\n\n"
                "Puede abrir el archivo en Excel para su consulta o cotejo."
            )
        except Exception as e:
            QMessageBox.critical(self, "Error al Exportar", f"No se pudo guardar el archivo Excel:\n{e}")

    def _on_v_modo_changed(self):
        """Actualiza el texto descriptivo del banner y limpia la selección previa."""
        modo = self.cmb_v_modo.currentData()
        self._v_parsed_records = []
        self.btn_v_ejecutar.setEnabled(False)
        self.btn_v_exportar_datos.setEnabled(False)
        self.lbl_vkpi_archivo.setText("Archivo: Ninguno")
        self.lbl_vkpi_total_arch.setText("En Archivo: 0")
        self.lbl_vkpi_nuevos.setText("Nuevos / Vinculables: 0")
        self.lbl_vkpi_duplicados.setText("Duplicados / Sin Coincidencia: 0")

        if hasattr(self, "lbl_v_desc") and self.lbl_v_desc is not None:
            if modo == "RPP":
                self.lbl_v_desc.setText(
                    "<b>2. Vinculación de Folios Electrónicos del RPP:</b><br/>"
                    "Cargue la relación de folios registrales para las viviendas ya existentes en el régimen.<br/>"
                    "• Relaciona cada registro mediante <b>EDIFICIO</b> e <b>INTERIOR</b>.<br/>"
                    "• Al asignar el folio RPP, la vivienda avanza automáticamente al estado <b>EN_INDIVIDUALIZACION</b>."
                )
            elif modo == "CATASTRO":
                self.lbl_v_desc.setText(
                    "<b>3. Vinculación de Cédula Catastral y Padrón Municipal:</b><br/>"
                    "Cargue las claves catastrales y cuentas prediales asignadas por Catastro Cancún.<br/>"
                    "• Enlaza mediante <b>FOLIO_ELECTRONICO</b> o por <b>(EDIFICIO, INTERIOR)</b>.<br/>"
                    "• Al completar los datos catastrales, la vivienda transiciona a <b>LISTA_PARA_VENTA</b>."
                )
            else:
                self.lbl_v_desc.setText(
                    "<b>1. Carga Inicial de Viviendas (Física):</b><br/>"
                    "Importe masivamente las viviendas/espacios privativos autorizados para el régimen.<br/>"
                    "• Se dan de alta en estado <b>PROYECTADA</b> (con interior, edificio, nivel, destino y espacio).<br/>"
                    "• No requieren folios RPP ni claves catastrales en esta etapa."
                )

        # Regla de Integridad Empresarial: Si el régimen ya cuenta con el 100% de UPEs cargadas,
        # es ilógico permitir Carga Inicial física. Conmutar a Vinculación RPP.
        if modo == "INICIAL":
            bd_count = getattr(self, "_v_bd_count", 0)
            tot_upes = getattr(self, "_v_tot_upes", 0)
            reg_id = self.cmb_v_regimen.currentData()
            if reg_id and tot_upes > 0 and bd_count >= tot_upes:
                QMessageBox.information(
                    self, "Régimen Completo al 100%",
                    f"El Régimen seleccionado ya cuenta con la totalidad de sus viviendas físicas cargadas "
                    f"({bd_count} de {tot_upes} UPEs autorizadas).\n\n"
                    "No es necesario realizar Carga Inicial. El sistema conmutará automáticamente a "
                    "'2. Vincular Folios RPP' para continuar con la individualización registral."
                )
                idx_rpp = self.cmb_v_modo.findData("RPP")
                if idx_rpp != -1:
                    self.cmb_v_modo.blockSignals(True)
                    self.cmb_v_modo.setCurrentIndex(idx_rpp)
                    self.cmb_v_modo.blockSignals(False)
                    self._on_v_modo_changed()
                return

        # Restaurar la tabla y métricas de forma instantánea sin disparar consultas ni parpadeo visual
        reg_id = self.cmb_v_regimen.currentData()
        if reg_id:
            cached_viviendas = getattr(self, "_v_db_viviendas", None)
            if cached_viviendas:
                # Restauración fluida desde la memoria (0 ms, sin parpadeo)
                self._render_viviendas_table_from_cache(cached_viviendas)
                if hasattr(self, "kpi_alertas_frame"):
                    self.lbl_vkpi_alertas.setText("Con Advertencia: 0")
                    self.kpi_alertas_frame.hide()
                if hasattr(self, "kpi_faltantes_frame"):
                    self.lbl_vkpi_faltantes.setText("Faltantes: 0")
                    self.kpi_faltantes_frame.hide()
                self._highlight_kpi_active("TODOS")
                self.btn_v_exportar_datos.setEnabled(len(cached_viviendas) > 0)
                self.action_state_changed.emit()
            else:
                self._on_v_regimen_changed()

    def _on_v_descargar_plantilla(self):
        """Genera una plantilla Excel (.xlsx) estandarizada según el modo de operación seleccionado."""
        des_id = self.cmb_v_desarrollo.currentData()
        reg_id = self.cmb_v_regimen.currentData()
        modo = self.cmb_v_modo.currentData()

        if not des_id or not reg_id:
            QMessageBox.warning(self, "Régimen Requerido", "Por favor seleccione un desarrollo y un régimen para generar la plantilla.")
            return

        nombre_des = "".join(c for c in self.cmb_v_desarrollo.currentText() if c.isalnum() or c in (" ", "_", "-")).strip()
        nombre_reg = "".join(c for c in self.cmb_v_regimen.currentText() if c.isalnum() or c in (" ", "_", "-")).strip()
        nombre_sugerido = f"Plantilla_{modo}_{nombre_des}_{nombre_reg}.xlsx"

        file_path, _ = QFileDialog.getSaveFileName(
            self, f"Guardar Plantilla Excel ({modo})", nombre_sugerido, "Excel Files (*.xlsx)"
        )
        if not file_path:
            return

        try:
            wb = openpyxl.Workbook()
            ws = wb.active

            if modo == "RPP":
                ws.title = "Folios_RPP"
                headers = ["No. Oficial", "Interior", "Edificio", "Nivel", "Folio", "Fecha"]
                sample_rows = [
                    ["UPE-001", "101", "TORRE A", "NIVEL 1", "RPP-CUN-0012345", "2026-08-27"],
                    ["UPE-002", "102", "TORRE A", "NIVEL 1", "RPP-CUN-0012346", "2026-08-27"],
                    ["UPE-003", "103", "TORRE A", "NIVEL 1", "RPP-CUN-0012347", "2026-08-27"]
                ]
            elif modo == "CATASTRO":
                ws.title = "Catastro_Cedula"
                headers = ["No. Oficial", "Interior", "Edificio", "Nivel", "Folio", "Cedula Catastral", "Fecha"]
                sample_rows = [
                    ["UPE-001", "101", "TORRE A", "NIVEL 1", "RPP-CUN-0012345", "101-01-001-001-001", "2026-08-27"],
                    ["UPE-002", "102", "TORRE A", "NIVEL 1", "RPP-CUN-0012346", "101-01-001-001-002", "2026-08-27"],
                    ["UPE-003", "103", "TORRE A", "NIVEL 1", "RPP-CUN-0012347", "101-01-001-001-003", "2026-08-27"]
                ]
            else:
                ws.title = "Viviendas_Inicial"
                headers = [
                    "NUMERO_OFICIAL", "INTERIOR", "EDIFICIO", "NIVEL", "CALLE",
                    "DESTINO", "ESPACIO", "DESCRIPCION"
                ]
                sample_rows = [
                    ["UPE-001", "101", "TORRE A", "NIVEL 1", "AV. HUAYACAN", "HABITACIONAL", "VIVIENDA", "Depto Tipo A 2 Recámaras"],
                    ["UPE-002", "102", "TORRE A", "NIVEL 1", "AV. HUAYACAN", "HABITACIONAL", "VIVIENDA", "Depto Tipo B 3 Recámaras"],
                    ["EST-001", "E-01", "TORRE A", "PLANTA BAJA", "AV. HUAYACAN", "SERVICIO", "AREAS COMUNES", "Cajón Privativo"]
                ]

            ws.append(headers)

            # Estilos de encabezado
            header_fill = PatternFill(start_color="1E293B", end_color="1E293B", fill_type="solid")
            header_font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
            border_thin = Border(
                left=Side(style='thin', color='CBD5E1'),
                right=Side(style='thin', color='CBD5E1'),
                top=Side(style='thin', color='CBD5E1'),
                bottom=Side(style='thin', color='CBD5E1')
            )

            for col_num in range(1, len(headers) + 1):
                cell = ws.cell(row=1, column=col_num)
                cell.fill = header_fill
                cell.font = header_font
                cell.alignment = Alignment(horizontal="center", vertical="center")
                cell.border = border_thin

            for row in sample_rows:
                ws.append(row)

            # Autoancho
            for col in ws.columns:
                max_len = max(len(str(cell.value or '')) for cell in col)
                col_letter = get_column_letter(col[0].column)
                ws.column_dimensions[col_letter].width = max(max_len + 5, 12)

            wb.save(file_path)
            QMessageBox.information(
                self, "Plantilla Generada",
                f"Plantilla Excel generada con éxito en:\n{file_path}\n\n"
                "Puede completarla con los predios y cargarla mediante 'Cargar Archivo'."
            )
        except Exception as e:
            QMessageBox.critical(self, "Error", f"No se pudo generar la plantilla: {e}")

    def _on_v_seleccionar_archivo(self):
        """Abre el diálogo de archivo y ejecuta la lectura y validación en segundo plano sin congelar la UI."""
        des_id = self.cmb_v_desarrollo.currentData()
        reg_id = self.cmb_v_regimen.currentData()
        modo = self.cmb_v_modo.currentData()

        if not des_id:
            QMessageBox.warning(self, "Desarrollo Requerido", "Por favor seleccione un Desarrollo antes de cargar el archivo.")
            return

        if not reg_id:
            QMessageBox.warning(
                self, "Régimen Requerido",
                "Por favor seleccione un Régimen en Condominio antes de cargar el archivo.\n\n"
                "Todo padrón de viviendas debe estar asociado a su régimen correspondiente para validar la capacidad de UPEs."
            )
            return

        if not modo:
            QMessageBox.warning(
                self, "Operación Requerida",
                "Por favor seleccione la Operación a realizar (Carga Inicial, Vincular Folios RPP o Catastro)."
            )
            return

        # Bloqueo estricto para Carga Inicial si el régimen ya está al 100%
        if modo == "INICIAL":
            bd_count = getattr(self, "_v_bd_count", 0)
            tot_upes = getattr(self, "_v_tot_upes", 0)
            if reg_id and tot_upes > 0 and bd_count >= tot_upes:
                QMessageBox.warning(
                    self, "Régimen Completo al 100%",
                    f"El Régimen seleccionado ya cuenta con la totalidad de sus viviendas cargadas "
                    f"({bd_count} de {tot_upes} UPEs autorizadas).\n\n"
                    "No es posible realizar Carga Inicial para este régimen. "
                    "Por favor seleccione la operación '2. Vincular Folios RPP' o '3. Vincular CC / Catastro'."
                )
                return

        filtro = "Todos los Soportados (*.xlsx *.xls *.pdf);;Constancias RPP o Catastro PDF (*.pdf);;Plantillas Excel (*.xlsx *.xls)"
        file_path, _ = QFileDialog.getOpenFileName(
            self, "Seleccionar Archivo (Excel o Constancia / Oficio PDF)", "", filtro
        )
        if not file_path:
            return

        is_pdf = file_path.lower().endswith(".pdf")
        if is_pdf and modo not in ("RPP", "CATASTRO"):
            QMessageBox.warning(
                self, "Modo de Operación Requerido",
                "Los archivos PDF son válidos para '2. Vincular Folios RPP' (Constancia de Resolución)\n"
                "o '3. Vincular CC / Catastro' (Oficio de Alta de Predio).\n\n"
                "Para '1. Carga Inicial' utilice la plantilla Excel."
            )
            return

        # Deshabilitar botones operativos mientras se analiza el archivo
        self.btn_v_seleccionar.setEnabled(False)
        self.btn_v_ejecutar.setEnabled(False)
        self.btn_v_exportar_datos.setEnabled(False)

        # Mostrar indicador de carga no bloqueante
        loading_msg = "Extrayendo datos del PDF (OCR)..." if is_pdf else "Analizando archivo y validando datos..."
        self._v_loading_dialog = GLLoadingDialog(loading_msg, self)
        self._v_loading_dialog.show()

        # Iniciar worker en background
        self._v_analyze_worker = ViviendasFileAnalyzeWorker(
            api_client=self.api_client,
            db_connector=self.db_connector,
            file_path=file_path,
            desarrollo_id=des_id,
            regimen_id=reg_id,
            modo=modo
        )
        self._v_analyze_worker.progress_updated.connect(self._on_v_file_progress_updated)
        self._v_analyze_worker.finished.connect(self._on_v_file_analyze_finished)
        self._v_analyze_worker.start()

    def _on_v_file_progress_updated(self, current: int, total: int, msg: str):
        """Actualiza el indicador numérico incremental dentro del load spinner."""
        if self._v_loading_dialog:
            self._v_loading_dialog.set_progress(current, total, msg)

    def _on_v_file_analyze_finished(self, success: bool, result_data: dict, error_msg: str):
        """Maneja el resultado del análisis asíncrono y actualiza la tabla y métricas de forma fluida."""
        if self._v_loading_dialog:
            self._v_loading_dialog.close()
            self._v_loading_dialog = None

        self.btn_v_seleccionar.setEnabled(True)

        if not success:
            QMessageBox.critical(self, "Error al Analizar Archivo", error_msg)
            return

        records = result_data.get("records", [])
        processed = result_data.get("processed_records", [])
        faltantes_bd = result_data.get("faltantes_bd", [])
        file_name = result_data.get("file_name", "")
        nuevos_cnt = result_data.get("nuevos_cnt", 0)
        duplicados_cnt = result_data.get("duplicados_cnt", 0)
        faltantes_cnt = len(faltantes_bd)
        elapsed_time = result_data.get("elapsed_time", 0.0)
        modo = self.cmb_v_modo.currentData()
        reg_id = self.cmb_v_regimen.currentData()
        reg_label = self.cmb_v_regimen.currentText() if reg_id else ""

        self._v_parsed_records = records
        self.lbl_vkpi_archivo.setText(f"Archivo: {file_name}")
        self.lbl_vkpi_total_arch.setText(f"En Archivo: {len(records)}")

        # Construir lista consolidada de filas para la tabla (Archivo + Faltantes en BD)
        table_rows = list(processed)
        for fb in faltantes_bd:
            table_rows.append({
                "data": fb,
                "status": "FALTANTE EN ARCHIVO"
            })

        self.tbl_viviendas.setUpdatesEnabled(False)
        try:
            self.tbl_viviendas.setRowCount(len(table_rows))
            for idx, item in enumerate(table_rows):
                r = item["data"]
                st_text = item["status"]

                item_st = QTableWidgetItem(st_text)
                if st_text in ("NUEVO", "VINCULABLE"):
                    item_st.setForeground(QColor("#16A34A"))
                elif st_text == "FALTANTE EN ARCHIVO":
                    item_st.setForeground(QColor("#D97706"))  # Ámbar de advertencia
                else:
                    item_st.setForeground(QColor("#DC2626"))

                self.tbl_viviendas.setItem(idx, 0, item_st)
                self.tbl_viviendas.setItem(idx, 1, QTableWidgetItem(reg_label or str(r.get("regimen_nombre") or "")))
                self.tbl_viviendas.setItem(idx, 2, QTableWidgetItem(str(r.get("numero_oficial") or "")))
                self.tbl_viviendas.setItem(idx, 3, QTableWidgetItem(str(r.get("interior") or "")))
                self.tbl_viviendas.setItem(idx, 4, QTableWidgetItem(str(r.get("edificio") or "")))
                self.tbl_viviendas.setItem(idx, 5, QTableWidgetItem(str(r.get("nivel") or "")))
                self.tbl_viviendas.setItem(idx, 6, QTableWidgetItem(str(r.get("destino") or "HABITACIONAL")))
                self.tbl_viviendas.setItem(idx, 7, QTableWidgetItem(str(r.get("espacio") or "VIVIENDA")))
                self.tbl_viviendas.setItem(idx, 8, QTableWidgetItem(str(r.get("folio_electronico") or "")))
                self.tbl_viviendas.setItem(idx, 9, QTableWidgetItem(str(r.get("clave_catastral") or "")))
                self.tbl_viviendas.setItem(idx, 10, QTableWidgetItem(str(r.get("padron") or "")))
                self.tbl_viviendas.setItem(idx, 11, QTableWidgetItem(str(r.get("calle") or "")))
                self.tbl_viviendas.setItem(idx, 12, QTableWidgetItem(str(r.get("descripcion") or "")))
        finally:
            self.tbl_viviendas.setUpdatesEnabled(True)

        self._v_nuevos_cnt = nuevos_cnt
        self._v_duplicados_cnt = duplicados_cnt
        self._v_duplicate_folios = result_data.get("duplicate_folios", [])
        self._v_duplicate_claves = result_data.get("duplicate_claves", [])

        if modo in ("RPP", "CATASTRO"):
            self.lbl_vkpi_nuevos.setText(f"Vinculables: {nuevos_cnt}")
            self.lbl_vkpi_duplicados.setText(f"Sin Coincidencia: {duplicados_cnt}")
            self.btn_v_ejecutar.setEnabled(nuevos_cnt > 0)
        else:
            self.lbl_vkpi_nuevos.setText(f"Nuevos: {nuevos_cnt}")
            self.lbl_vkpi_duplicados.setText(f"Duplicados: {duplicados_cnt}")
            self.btn_v_ejecutar.setEnabled(nuevos_cnt > 0)

        # Contabilizar registros con advertencia de inconsistencia en el documento
        alertas_cnt = sum(1 for row_it in table_rows if "⚠️" in str(row_it.get("data", {}).get("descripcion") or "") or "ERROR EN TEXTO" in str(row_it.get("data", {}).get("descripcion") or ""))
        if alertas_cnt > 0:
            self.lbl_vkpi_alertas.setText(f"Con Advertencia: {alertas_cnt}")
            self.kpi_alertas_frame.show()
        else:
            self.kpi_alertas_frame.hide()

        # Actualizar KPI de Faltantes y visibilidad
        if faltantes_cnt > 0:
            self.lbl_vkpi_faltantes.setText(f"Faltantes en Archivo: {faltantes_cnt}")
            self.kpi_faltantes_frame.show()
        else:
            self.kpi_faltantes_frame.hide()

        # Poblar selector de filtro de estado
        self.cmb_v_filtro_estado.blockSignals(True)
        self.cmb_v_filtro_estado.clear()
        self.cmb_v_filtro_estado.addItem(f"Mostrar: Todos ({len(table_rows)})", "TODOS")
        if modo in ("RPP", "CATASTRO"):
            self.cmb_v_filtro_estado.addItem(f"Solo Vinculables ({nuevos_cnt})", "VINCULABLE")
            if duplicados_cnt > 0:
                self.cmb_v_filtro_estado.addItem(f"Solo Sin Coincidencia ({duplicados_cnt})", "NO ENCONTRADO")
            if faltantes_cnt > 0:
                self.cmb_v_filtro_estado.addItem(f"Solo Faltantes en Archivo ({faltantes_cnt})", "FALTANTE EN ARCHIVO")
        else:
            self.cmb_v_filtro_estado.addItem(f"Solo Nuevos ({nuevos_cnt})", "NUEVO")
            if duplicados_cnt > 0:
                self.cmb_v_filtro_estado.addItem(f"Solo Duplicados ({duplicados_cnt})", "DUPLICADO")
        self.cmb_v_filtro_estado.setCurrentIndex(0)
        self.cmb_v_filtro_estado.blockSignals(False)

        self.btn_v_exportar_datos.setEnabled(len(records) > 0)
        self.action_state_changed.emit()

        # Mensaje de confirmación de análisis con tiempo de procesamiento y resumen de registros
        tiempo_str = f"{int(elapsed_time // 60)} min {int(elapsed_time % 60)} seg" if elapsed_time >= 60 else f"{elapsed_time:.1f} seg"
        resumen_txt = (
            f"Archivo analizado exitosamente en {tiempo_str}.\n\n"
            f"• Total en archivo: {len(records)} registros\n"
        )
        if modo in ("RPP", "CATASTRO"):
            resumen_txt += (
                f"• Vinculables con el régimen: {nuevos_cnt}\n"
                f"• Sin coincidencia: {duplicados_cnt}\n"
            )
            if faltantes_cnt > 0:
                resumen_txt += f"• Faltantes en archivo (en BD): {faltantes_cnt}\n"
        else:
            resumen_txt += (
                f"• Nuevos para crear: {nuevos_cnt}\n"
                f"• Duplicados omitidos: {duplicados_cnt}\n"
            )

        QMessageBox.information(self, "Análisis de Archivo Completado", resumen_txt)

    def _on_v_ejecutar_importacion(self):
        """Ejecuta la importación en un hilo background mediante ViviendasImportWorker mostrando indicador visual."""
        des_id = self.cmb_v_desarrollo.currentData()
        reg_id = self.cmb_v_regimen.currentData()

        if not des_id or not self._v_parsed_records:
            return

        total_recs = len(self._v_parsed_records)
        target_cnt = getattr(self, "_v_nuevos_cnt", total_recs)
        sin_coincidencia = getattr(self, "_v_duplicados_cnt", 0)
        dup_folios = getattr(self, "_v_duplicate_folios", [])
        dup_claves = getattr(self, "_v_duplicate_claves", [])
        reg_txt = self.cmb_v_regimen.currentText() if reg_id else "Directo a Desarrollo"
        des_txt = self.cmb_v_desarrollo.currentText()

        advertencia_dups_str = ""
        if dup_folios:
            f_ej = ", ".join([d[0] for d in dup_folios[:3]])
            advertencia_dups_str += (
                f"<br/><b style='color: #DC2626;'>⚠️ ADVERTENCIA: Hay {len(dup_folios)} Folio(s) RPP repetidos en varias filas del archivo "
                f"({f_ej}{'...' if len(dup_folios) > 3 else ''}).</b>"
            )
        if dup_claves:
            c_ej = ", ".join([d[0] for d in dup_claves[:3]])
            advertencia_dups_str += (
                f"<br/><b style='color: #DC2626;'>⚠️ ADVERTENCIA: Hay {len(dup_claves)} Clave(s) Catastral(es) repetidas en varias filas del archivo "
                f"({c_ej}{'...' if len(dup_claves) > 3 else ''}).</b>"
            )

        modo = self.cmb_v_modo.currentData()
        if modo == "RPP":
            modo_titulo = "Vinculación de Folios RPP"
            modo_desc = f"¿Está seguro de vincular los <b>Folios Electrónicos RPP</b> para las <b>{target_cnt}</b> viviendas detectadas?{advertencia_dups_str}"
            loading_txt = f"Vinculando {target_cnt} Folios RPP...\nPor favor espere."
            detalle_procesar = f"<b>Total a vincular:</b> {target_cnt} viviendas (de {total_recs} en archivo)"
            if sin_coincidencia > 0:
                detalle_procesar += f"<br/><b>Sin coincidencia:</b> {sin_coincidencia} registros omitidos"
        elif modo == "CATASTRO":
            modo_titulo = "Vinculación Catastral y Predial"
            modo_desc = f"¿Está seguro de vincular las <b>Claves Catastrales y Padrones</b> para las <b>{target_cnt}</b> viviendas detectadas?{advertencia_dups_str}"
            loading_txt = f"Vinculando Catastro y Padrón para {target_cnt} viviendas...\nPor favor espere."
            detalle_procesar = f"<b>Total a vincular:</b> {target_cnt} viviendas (de {total_recs} en archivo)"
            if sin_coincidencia > 0:
                detalle_procesar += f"<br/><b>Sin coincidencia:</b> {sin_coincidencia} registros omitidos"
        else:
            modo_titulo = "Carga Inicial de Viviendas"
            tot_regimen = getattr(self, "_v_tot_upes", 0)
            advertencia_parcial = ""
            if tot_regimen > 0 and target_cnt < tot_regimen:
                faltantes_dif = tot_regimen - target_cnt
                advertencia_parcial = (
                    f"<br/><br/><b style='color: #D97706;'>⚠️ ADVERTENCIA DE CARGA PARCIAL:</b><br/>"
                    f"El archivo contiene <b>{target_cnt}</b> UPEs, lo cual es MENOR a la capacidad autorizada "
                    f"del Régimen (<b>{tot_regimen}</b> UPEs). Faltarían <b>{faltantes_dif}</b> unidades por dar de alta.<br/>"
                    f"¿Desea continuar con esta carga parcial?"
                )

            modo_desc = f"¿Está seguro de aplicar la <b>Carga Inicial</b> de <b>{target_cnt}</b> viviendas nuevas en estado <i>PROYECTADA</i>?{advertencia_parcial}"
            loading_txt = f"Importando {target_cnt} viviendas nuevas...\nPor favor espere."
            detalle_procesar = f"<b>Nuevas a crear:</b> {target_cnt} viviendas (de {total_recs} en archivo)"
            if sin_coincidencia > 0:
                detalle_procesar += f"<br/><b>Duplicadas:</b> {sin_coincidencia} registros omitidos"

        msg_confirm = (
            f"{modo_desc}<br/><br/>"
            f"<b>Desarrollo:</b> {des_txt}<br/>"
            f"<b>Régimen:</b> {reg_txt}<br/>"
            f"{detalle_procesar}"
        )

        reply = QMessageBox.question(
            self, f"Confirmar {modo_titulo}",
            msg_confirm,
            QMessageBox.Yes | QMessageBox.No,
            QMessageBox.Yes
        )
        if reply != QMessageBox.Yes:
            return

        self.btn_v_ejecutar.setEnabled(False)
        self.btn_v_seleccionar.setEnabled(False)

        # Mostrar indicador de carga no bloqueante
        self._v_loading_dialog = GLLoadingDialog(loading_txt, self)
        self._v_loading_dialog.show()

        self._v_worker = ViviendasImportWorker(
            api_client=self.api_client,
            db_connector=self.db_connector,
            desarrollo_id=des_id,
            regimen_id=reg_id,
            records=self._v_parsed_records,
            usuario_id=self.current_user_id or 1,
            sesion_id=self.current_sesion_id,
            modo=modo
        )
        self._v_worker.finished.connect(self._on_v_worker_finished)
        self._v_worker.start()

    def _on_v_worker_finished(self, success: bool, result: dict, error_msg: str):
        # Cerrar el indicador de carga si está activo
        if self._v_loading_dialog:
            self._v_loading_dialog.close()
            self._v_loading_dialog = None

        self.btn_v_seleccionar.setEnabled(True)
        modo = self.cmb_v_modo.currentData()
        if success:
            el_t = result.get("elapsed_time", 0.0) if isinstance(result, dict) else 0.0
            t_txt = f" en {int(el_t // 60)} min {int(el_t % 60)} seg" if el_t >= 60 else f" en {el_t:.1f} seg" if el_t > 0 else ""

            if modo == "RPP":
                act = result.get("actualizados", 0)
                no_enc = result.get("no_encontrados", 0)
                QMessageBox.information(
                    self, "Vinculación RPP Exitosa",
                    f"Folios RPP vinculados exitosamente{t_txt}:\n\n"
                    f"• Viviendas actualizadas a 'EN_INDIVIDUALIZACION': {act}\n"
                    f"• Registros no encontrados en el régimen: {no_enc}"
                )
            elif modo == "CATASTRO":
                act = result.get("actualizados", 0)
                no_enc = result.get("no_encontrados", 0)
                QMessageBox.information(
                    self, "Vinculación Catastral Exitosa",
                    f"Claves Catastrales y Padrones vinculados exitosamente{t_txt}:\n\n"
                    f"• Viviendas actualizadas a 'LISTA_PARA_VENTA': {act}\n"
                    f"• Registros no encontrados en el régimen: {no_enc}"
                )
            else:
                nuevos = result.get("nuevos_insertados", 0)
                dups = result.get("omitidos_duplicados", 0)
                QMessageBox.information(
                    self, "Carga Inicial Exitosa",
                    f"Proceso de importación inicial completado exitosamente{t_txt}:\n\n"
                    f"• Viviendas creadas en 'PROYECTADA': {nuevos}\n"
                    f"• Omitidas por duplicidad: {dups}"
                )

            self._v_parsed_records = []
            self.tbl_viviendas.setUpdatesEnabled(False)
            try:
                self._on_v_regimen_changed()
            finally:
                self.tbl_viviendas.setUpdatesEnabled(True)
        else:
            QMessageBox.critical(self, "Error en la Operación", f"No se pudo completar la operación:\n{error_msg}")
            self.btn_v_ejecutar.setEnabled(True)

